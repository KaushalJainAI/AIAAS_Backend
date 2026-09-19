"""
The office tools: `render_deck`, `render_workbook`, `render_document`.

Every renderer is tested by **reading the file back with a real reader**
(python-pptx, openpyxl, python-docx), because the failure that matters is a
file that saves without error and does not open — or opens with the numbers
typed in where formulas should be. A test that only checks bytes were written
passes for all of those.
"""
from __future__ import annotations

import io
import json
import shutil
import tempfile

import docx
import openpyxl
from asgiref.sync import async_to_sync
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from pptx import Presentation

from chat.tools import execute_tool, get_available_tools
from chat.tools.office import OFFICE_TOOLS, _CHART_SCHEMA, deck, document, workbook
from chat.tools.office.spec import SpecError
from chat.tools.registry import get as registered

User = get_user_model()

CHART = {
    'kind': 'column', 'title': 'Revenue by region', 'y_label': 'INR lakh',
    'series': [
        {'name': '2025', 'points': [{'x': 'North', 'y': 120}, {'x': 'South', 'y': 95}]},
        {'name': '2026', 'points': [{'x': 'North', 'y': 150}, {'x': 'South', 'y': None}]},
    ],
}

#: One slide of every layout, so bounds and round-trip tests cover them all.
EVERY_LAYOUT = [
    {'layout': 'title', 'title': 'EV market in India', 'subtitle': '2026 outlook'},
    {'layout': 'section', 'title': 'Where we are'},
    {'layout': 'bullets', 'title': 'What changed',
     'bullets': ['**Subsidies** moved to makers', '- FAME III replaced FAME II', 'Charging doubled'],
     'notes': 'Stress the subsidy shift.'},
    {'layout': 'two_column', 'title': 'Us vs them',
     'left': {'heading': 'Strengths', 'bullets': ['Price']},
     'right': {'heading': 'Gaps', 'bullets': ['Range']}},
    {'layout': 'chart', 'title': 'Revenue', 'chart': CHART, 'caption': 'Source: ledger'},
    {'layout': 'table', 'title': 'Models', 'columns': ['Model', 'Range'],
     'rows': [['Ather', '150 km'], ['Ola', '190 km']]},
    {'layout': 'quote', 'quote': 'Electrify early.', 'attribution': 'An analyst'},
    {'layout': 'stats', 'title': 'Numbers', 'stats': [
        {'value': '₹4.2 Cr', 'label': 'Revenue'}, {'value': '38%', 'label': 'Share'},
        {'value': '1.2M', 'label': 'Units'}, {'value': '14%', 'label': 'Cheaper cells'}]},
    {'layout': 'closing', 'title': 'Thank you'},
]


def _deck(slides=EVERY_LAYOUT, **extra):
    return deck.validate({'slides': slides, **extra})


# ─────────────────────────────────────────────────────────────────────────
# Decks
# ─────────────────────────────────────────────────────────────────────────

class DeckTests(SimpleTestCase):
    def test_every_layout_renders_in_every_theme_and_opens(self):
        for theme in ('clean', 'bold', 'dark'):
            with self.subTest(theme=theme):
                data = deck.render(_deck(theme=theme), {})
                prs = Presentation(io.BytesIO(data))
                self.assertEqual(len(prs.slides), len(EVERY_LAYOUT))

    def test_nothing_is_drawn_off_the_slide(self):
        # The layout is ours, so text off the edge is our bug, not the model's.
        prs = Presentation(io.BytesIO(deck.render(_deck(), {})))
        for n, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                with self.subTest(slide=n, shape=shape.name):
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                    self.assertLessEqual(shape.top + shape.height, prs.slide_height)

    def test_charts_are_native_and_keep_their_numbers(self):
        prs = Presentation(io.BytesIO(deck.render(_deck(), {})))
        charts = [s.chart for sl in prs.slides for s in sl.shapes if s.has_chart]
        self.assertEqual(len(charts), 1)
        plot = charts[0].plots[0]
        self.assertEqual(list(plot.categories), ['North', 'South'])
        self.assertEqual([s.name for s in plot.series], ['2025', '2026'])
        # A missing measurement is a gap, never a zero.
        self.assertEqual(list(plot.series[1].values), [150.0, None])

    def test_notes_and_title_survive(self):
        prs = Presentation(io.BytesIO(deck.render(_deck(title='EV deck'), {})))
        self.assertEqual(prs.core_properties.title, 'EV deck')
        self.assertIn('subsidy shift', prs.slides[2].notes_slide.notes_text_frame.text)

    def test_sub_points_and_bold_are_parsed(self):
        spec = _deck()
        bullets = spec['slides'][2]['bullets']
        self.assertEqual([b['level'] for b in bullets], [0, 1, 0])
        self.assertEqual(bullets[1]['text'], 'FAME III replaced FAME II')

    def test_refusals_name_the_slide_and_the_fix(self):
        cases = [
            ({'layout': 'bullets', 'title': 't', 'bullets': ['x'] * 7}, 'limit is 6'),
            ({'layout': 'bullets', 'title': 'x' * 200, 'bullets': ['a']}, 'Shorten'),
            ({'layout': 'hologram', 'title': 't'}, 'layout must be one of'),
            ({'layout': 'bullets', 'bullets': ['a']}, 'title is required'),
            ({'layout': 'bullets', 'title': 't', 'bullets': ['- orphan']}, 'sub-point'),
            ({'layout': 'chart', 'chart': {**CHART, 'kind': 'pie'}}, 'one series'),
            ({'layout': 'chart', 'chart': {'kind': 'scatter', 'title': 's', 'series': [
                {'name': 'a', 'points': [{'x': 'big', 'y': 1}]}]}}, 'numeric x'),
            ({'layout': 'stats', 'title': 't', 'stats': [{'value': 'x', 'label': 'y'}] * 5},
             'limit is 4'),
        ]
        for slide, message in cases:
            with self.subTest(message=message):
                with self.assertRaisesRegex(SpecError, message) as ctx:
                    _deck([EVERY_LAYOUT[0], slide])
                self.assertIn('Slide 2', str(ctx.exception))

    def test_too_many_slides_is_refused(self):
        with self.assertRaisesRegex(SpecError, 'limit is 40'):
            _deck([EVERY_LAYOUT[0]] * 41)

    def test_the_stored_preview_carries_the_theme_tokens(self):
        preview = deck.preview(_deck(theme='dark'))
        self.assertEqual(preview['kind'], 'deck')
        self.assertEqual(preview['theme']['background'], '141414')
        self.assertEqual(len(preview['slides']), len(EVERY_LAYOUT))

    def test_the_extract_holds_what_the_slides_say(self):
        text = deck.extract_text(_deck())
        for fragment in ('EV market in India', 'FAME III', 'Ather | 150 km',
                         '₹4.2 Cr — Revenue', 'Notes: Stress'):
            self.assertIn(fragment, text)

    def test_chart_slides_accept_exactly_render_charts_shape(self):
        # One validator: the tool schema for a chart on a slide *is*
        # render_chart's, so the two cannot drift.
        self.assertEqual(
            {k: v for k, v in _CHART_SCHEMA.items() if k != 'description'},
            {k: v for k, v in registered('render_chart').schema['function']['parameters'].items()
             if k != 'description'},
        )


# ─────────────────────────────────────────────────────────────────────────
# Workbooks
# ─────────────────────────────────────────────────────────────────────────

SALES = {
    'name': 'Sales',
    'columns': [
        {'header': 'Region'},
        {'header': 'Units', 'type': 'integer'},
        {'header': 'Price', 'type': 'currency', 'currency': 'INR'},
        {'header': 'Revenue', 'type': 'currency', 'currency': 'INR'},
        {'header': 'Margin', 'type': 'percent'},
        {'header': 'Date', 'type': 'date'},
    ],
    'rows': [
        ['North', 120, 1500, '=B{r}*C{r}', 0.21, '2026-01-31'],
        ['South', 95, '1,450', '=B{r}*C{r}', '19%', '2026-02-28'],
        {'Region': 'West', 'Units': 140, 'Price': 1600, 'Revenue': '=B{r}*C{r}'},
    ],
    'totals': True,
    'chart': {'kind': 'column', 'title': 'Revenue by region', 'x': 'Region', 'y': ['Revenue']},
}


def _book(*sheets):
    spec = workbook.validate({'sheets': list(sheets) or [SALES]})
    data, warnings = workbook.render(spec)
    return openpyxl.load_workbook(io.BytesIO(data)), warnings, spec


class WorkbookTests(SimpleTestCase):
    def test_formulas_stay_formulas_and_r_is_this_row(self):
        wb, _, _ = _book()
        ws = wb['Sales']
        self.assertEqual(ws['D2'].value, '=B2*C2')
        self.assertEqual(ws['D4'].value, '=B4*C4')

    def test_totals_are_a_sum_not_a_typed_number(self):
        ws = _book()[0]['Sales']
        self.assertEqual(ws['A5'].value, 'Total')
        self.assertEqual(ws['B5'].value, '=SUM(B2:B4)')
        self.assertEqual(ws['D5'].value, '=SUM(D2:D4)')
        # Percentages do not add up to anything, so they are not totalled.
        self.assertIsNone(ws['E5'].value)

    def test_values_are_typed_by_their_column(self):
        ws = _book()[0]['Sales']
        self.assertEqual(ws['C3'].value, 1450)          # "1,450" in a currency column
        self.assertAlmostEqual(ws['E3'].value, 0.19)    # "19%" in a percent column
        self.assertEqual(ws['F2'].value.year, 2026)     # a real date, not text
        self.assertIn('₹', ws['C2'].number_format)
        self.assertEqual(ws['E2'].number_format, '0.0%')

    def test_header_is_frozen_filtered_and_charted(self):
        ws = _book()[0]['Sales']
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:F4')
        self.assertEqual(len(ws._charts), 1)

    def test_a_non_number_in_a_number_column_is_kept_and_reported(self):
        sheet = {'name': 'S', 'columns': [{'header': 'N', 'type': 'number'}], 'rows': [['n/a']]}
        wb, warnings, _ = _book(sheet)
        self.assertEqual(wb['S']['A2'].value, 'n/a')
        self.assertTrue(warnings)

    def test_refusals(self):
        cases = [
            ({**SALES, 'rows': [['a'] * 7]}, '7 values but there are 6'),
            ({**SALES, 'rows': [['a', 1, 1, '=WEBSERVICE("http://x")']]}, 'outside the'),
            ({**SALES, 'rows': [['a', 1, 1, '=cmd|"/c calc"!A1']]}, 'outside the'),
            ({**SALES, 'chart': {**SALES['chart'], 'y': ['Region']}}, 'needs numbers'),
            ({**SALES, 'chart': {**SALES['chart'], 'x': 'Nope'}}, 'not a column'),
            ({**SALES, 'totals': ['Margin']}, 'cannot be totalled'),
            ({**SALES, 'columns': ['A', 'a']}, 'share a header'),
        ]
        for sheet, message in cases:
            with self.subTest(message=message):
                with self.assertRaisesRegex(SpecError, message):
                    workbook.validate({'sheets': [sheet]})

    def test_the_preview_is_capped_and_says_so(self):
        rows = [[i] for i in range(workbook.PREVIEW_ROWS + 5)]
        spec = workbook.validate({'sheets': [
            {'name': 'Big', 'columns': [{'header': 'n', 'type': 'integer'}], 'rows': rows}]})
        sheet = workbook.preview(spec)['sheets'][0]
        self.assertEqual(len(sheet['rows']), workbook.PREVIEW_ROWS)
        self.assertEqual(sheet['row_count'], workbook.PREVIEW_ROWS + 5)


# ─────────────────────────────────────────────────────────────────────────
# Documents
# ─────────────────────────────────────────────────────────────────────────

class DocumentTests(SimpleTestCase):
    SPEC = {
        'title': 'Q3 review', 'subtitle': 'For the board',
        'blocks': [
            {'type': 'heading', 'text': 'Summary', 'level': 1},
            {'type': 'paragraph', 'text': 'Revenue grew **38%** while costs rose *9%*.'},
            {'type': 'bullets', 'items': ['One', 'Two']},
            {'type': 'numbered', 'items': ['First']},
            {'type': 'table', 'columns': ['A', 'B'], 'rows': [['1', '2']], 'caption': 'Table 1'},
            {'type': 'chart', 'chart': CHART},
            {'type': 'page_break'},
        ],
    }

    def _open(self, spec=None):
        data = document.render(document.validate(spec or self.SPEC), {})
        return docx.Document(io.BytesIO(data))

    def test_structure_survives(self):
        d = self._open()
        styles = [p.style.name for p in d.paragraphs if p.text]
        self.assertEqual(d.paragraphs[0].text, 'Q3 review')
        self.assertIn('Heading 1', styles)
        self.assertIn('List Bullet', styles)
        self.assertIn('List Number', styles)

    def test_inline_emphasis_becomes_runs(self):
        para = next(p for p in self._open().paragraphs if 'Revenue' in p.text)
        self.assertEqual(para.text, 'Revenue grew 38% while costs rose 9%.')
        self.assertTrue(any(r.bold and r.text == '38%' for r in para.runs))
        self.assertTrue(any(r.italic and r.text == '9%' for r in para.runs))

    def test_a_chart_block_is_saved_as_its_data(self):
        d = self._open()
        self.assertEqual(len(d.tables), 2)
        chart_table = d.tables[1]
        self.assertEqual([c.text for c in chart_table.rows[0].cells], ['Category', '2025', '2026'])
        self.assertEqual([c.text for c in chart_table.rows[2].cells], ['South', '95', '—'])

    def test_dark_is_not_a_document_theme(self):
        with self.assertRaisesRegex(SpecError, 'theme'):
            document.validate({**self.SPEC, 'theme': 'dark'})


# ─────────────────────────────────────────────────────────────────────────
# The tools: scope, paths, and what the model is told
# ─────────────────────────────────────────────────────────────────────────

class OfficeToolTests(TestCase):
    def setUp(self):
        from inference import vfs

        self._media = tempfile.mkdtemp()
        self._override = override_settings(MEDIA_ROOT=self._media)
        self._override.enable()
        self.user = User.objects.create_user('owner', 'owner@example.com', 'pw')
        self.scope = vfs.chat_scope(self.user)

    def tearDown(self):
        self._override.disable()
        shutil.rmtree(self._media, ignore_errors=True)

    def _call(self, name, args, scope='default'):
        ctx = {'user_id': self.user.id, 'file_scope': self.scope if scope == 'default' else scope}
        return json.loads(async_to_sync(execute_tool)(name, args, ctx))

    def test_a_bare_name_lands_in_the_writable_folder(self):
        out = self._call('render_deck', {'path': 'q3', 'slides': EVERY_LAYOUT[:2]})
        self.assertEqual(out['path'], '/Chat/q3.pptx')
        self.assertEqual(out['slides'], 2)
        self.assertIn('do not paste', out['rendered'])

    def test_no_path_is_named_after_the_content(self):
        out = self._call('render_workbook', {'sheets': [SALES]})
        self.assertEqual(out['path'], '/Chat/Sales.xlsx')
        self.assertEqual(out['sheets'], [{'name': 'Sales', 'rows': 3}])

    def test_a_second_render_is_renamed_and_says_so(self):
        self._call('render_document', {'path': 'r.docx', **DocumentTests.SPEC})
        out = self._call('render_document', {'path': 'r.docx', **DocumentTests.SPEC})
        self.assertEqual(out['path'], '/Chat/r (2).docx')
        self.assertIn('overwrite=true', out['note'])
        self.assertEqual(out['charts_as_tables'], 1)

    def test_the_wrong_extension_is_refused(self):
        out = self._call('render_deck', {'path': '/Chat/q3.docx', 'slides': EVERY_LAYOUT[:1]})
        self.assertIn('.pptx', out['error'])

    def test_a_spec_error_comes_back_as_an_error_not_a_crash(self):
        out = self._call('render_deck', {'slides': [{'layout': 'bullets', 'title': 't'}]})
        self.assertIn('Slide 1', out['error'])

    def test_no_scope_no_file(self):
        out = self._call('render_deck', {'slides': EVERY_LAYOUT[:1]}, scope=None)
        self.assertIn('no file workspace', out['error'])

    def test_an_image_slide_embeds_a_picture_from_the_users_files(self):
        from PIL import Image

        from inference import vfs

        buf = io.BytesIO()
        Image.new('RGB', (400, 200), 'navy').save(buf, format='PNG')
        vfs.write_binary(self.scope, '/Chat/chart.png', buf.getvalue())
        out = self._call('render_deck', {'path': 'pics', 'slides': [
            {'layout': 'image', 'title': 'Look', 'image': '/Chat/chart.png'}]})
        self.assertNotIn('error', out)
        from inference.models import Document

        with Document.objects.get(id=out['document_id']).file.open('rb') as fh:
            prs = Presentation(fh)
        pictures = [s for s in prs.slides[0].shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 1)

    def test_a_missing_image_is_an_error_naming_it(self):
        out = self._call('render_deck', {'slides': [
            {'layout': 'image', 'title': 'Look', 'image': '/Chat/nope.png'}]})
        self.assertIn('No such image', out['error'])


class OfficeAvailabilityTests(SimpleTestCase):
    def test_offered_only_with_a_file_scope(self):
        bare = {t['function']['name'] for t in async_to_sync(get_available_tools)(None)}
        scoped = {t['function']['name']
                  for t in async_to_sync(get_available_tools)(None, file_scope=object())}
        for name in OFFICE_TOOLS:
            self.assertNotIn(name, bare)
            self.assertIn(name, scoped)

    def test_they_create_rather_than_destroy(self):
        # Why they run without asking in chat: nothing they do is beyond the
        # user's own undo (renamed on collision, overwrite goes to the bin).
        for name in OFFICE_TOOLS:
            tool = registered(name)
            self.assertEqual(tool.effect, 'reversible')
            self.assertFalse(tool.sensitive)

    def test_the_office_grant_unlocks_them_and_needs_a_scope(self):
        from agents.agent.runtime import GRANT_TOOLS, AgentToolbox

        self.assertEqual(set(GRANT_TOOLS['office']), set(OFFICE_TOOLS))
        without = AgentToolbox(grants={'office': True}, user_id=1).allowed_names
        with_scope = AgentToolbox(grants={'office': True}, user_id=1,
                                  file_scope=object()).allowed_names
        self.assertFalse(set(OFFICE_TOOLS) & without)
        self.assertTrue(set(OFFICE_TOOLS) <= with_scope)
        # And it is not a back door to the rest of the file tools.
        self.assertNotIn('write_file', with_scope)

    def test_plan_mode_withholds_them(self):
        from agents.agent.runtime import AgentToolbox

        box = AgentToolbox(grants={'office': True}, user_id=1, file_scope=object(),
                           read_only=True)
        self.assertFalse(set(OFFICE_TOOLS) & box.allowed_names)
