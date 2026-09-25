"""
`.docx` from a spec: a title block and a list of typed blocks.

Blocks rather than markdown, for the reason the deck takes layouts: a
markdown string is a format the model *writes*, so every document would be
whatever the model remembered about markdown that day, and a malformed table
would render as pipes. Typed blocks are validated here and styled here.

`**bold**` and `*italic*` inside a paragraph are still honoured, because
emphasis is part of what a sentence says rather than of how the page looks.

Charts cannot be native in a Word file without a server-side renderer, and
there is deliberately no plotting library (a second chart design system is
what `render_chart` exists to avoid). So a `chart` block is rendered as its
data table with a caption saying so — the numbers survive, and the result
tells the model it happened rather than letting it believe the page has a
picture on it.
"""
from __future__ import annotations

import io
import re
from typing import Any

from docx import Document as new_docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..charts import ChartError, build_spec
from .spec import SpecError, choice, items, text
from .themes import THEMES, Theme

BLOCK_TYPES = ('heading', 'paragraph', 'bullets', 'numbered', 'table', 'quote',
               'image', 'chart', 'page_break')
#: Documents are read on white and printed, so the dark theme is not offered.
DOC_THEMES = ('clean', 'bold')

#: Paragraph and heading alignment.
ALIGNMENTS = ('left', 'center', 'right', 'justify')

MAX_BLOCKS = 300
TITLE_CHARS = 200
HEADING_CHARS = 200
PARAGRAPH_CHARS = 6000
LIST_ITEMS = 50
ITEM_CHARS = 1000
TABLE_COLS = 12
TABLE_ROWS = 200
CELL_CHARS = 500
CAPTION_CHARS = 300
RUN_TEXT_CHARS = 6000
LINK_CHARS = 500

_INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\s][^*]*?\*)')


def validate(args: dict, *, max_blocks: int = MAX_BLOCKS) -> dict:
    """The normalised spec, or `SpecError`. The knob
    (`render_document/render_pdf.maxBlocks`) narrows, never widens."""
    theme = choice(args.get('theme'), 'theme', DOC_THEMES, 'clean')
    title = text(args.get('title'), 'title', TITLE_CHARS, required=True)
    subtitle = text(args.get('subtitle'), 'subtitle', TITLE_CHARS)
    blocks = []
    for n, raw in enumerate(items(args.get('blocks'), 'blocks',
                                  max(1, min(max_blocks, MAX_BLOCKS)),
                                  required=True), 1):
        if not isinstance(raw, dict):
            raise SpecError(f'Block {n} must be an object with a type.')
        kind = choice(raw.get('type'), f'Block {n} type', BLOCK_TYPES, 'paragraph')
        try:
            blocks.append(_block(raw, kind))
        except SpecError as exc:
            raise SpecError(f'Block {n} ({kind}): {exc}') from None
    return {'title': title, 'subtitle': subtitle, 'theme': theme, 'blocks': blocks}


def _run(raw: Any, where: str) -> dict:
    """One inline run: text with its own bold/italic/underline/strike/code/link.

    The text is *not* stripped, unlike every other string field: edge spaces
    are what separate one run from the next, and stripping them glues runs
    together.
    """
    if not isinstance(raw, dict):
        raise SpecError(f'{where} must be an object with text.')
    value = raw.get('text')
    if not isinstance(value, str) or not value:
        raise SpecError(f'{where} text is required.')
    if len(value) > RUN_TEXT_CHARS:
        raise SpecError(f'{where} text has {len(value)} characters; the limit is {RUN_TEXT_CHARS}.')
    run: dict[str, Any] = {'text': value}
    for flag in ('bold', 'italic', 'underline', 'strike', 'code'):
        if flag in raw:
            if not isinstance(raw[flag], bool):
                raise SpecError(f'{where} {flag} must be true or false.')
            if raw[flag]:
                run[flag] = True
    if raw.get('link') not in (None, ''):
        run['link'] = text(raw.get('link'), f'{where} link', LINK_CHARS)
    return run


def _runs_or_text(raw: dict, where: str, limit: int) -> dict:
    """`runs` when given (validated), else the legacy marker `text`."""
    if raw.get('runs') is not None:
        runs = [ _run(r, f'{where} run {n}')
                 for n, r in enumerate(items(raw.get('runs'), f'{where} runs', LIST_ITEMS),
                                       1)]
        if not runs:
            raise SpecError(f'{where} runs must hold at least one run.')
        return {'runs': runs}
    return {'text': text(raw.get('text'), where, limit, required=True)}


def _align(raw: dict) -> str:
    align = str(raw.get('align') or 'left').lower()
    if align not in ALIGNMENTS:
        raise SpecError(f'align must be one of {list(ALIGNMENTS)}.')
    return align


def _item(raw: Any, n: int) -> Any:
    """A list item: a marker string, or an object with text/runs like a paragraph."""
    if isinstance(raw, str):
        return text(raw, 'item', ITEM_CHARS, required=True)
    if isinstance(raw, dict):
        if raw.get('runs') is not None:
            return {'runs': [_run(r, f'item {n} run {m}')
                             for m, r in enumerate(
                                 items(raw.get('runs'), f'item {n} runs', 20), 1)]}
        return {'text': text(raw.get('text'), 'item', ITEM_CHARS, required=True)}
    raise SpecError(f'item {n} must be text or an object with text.')


def _block(raw: dict, kind: str) -> dict:
    b: dict[str, Any] = {'type': kind}
    if kind == 'heading':
        level = raw.get('level', 1)
        b['level'] = level if level in (1, 2, 3) else 1
        b.update(_runs_or_text(raw, 'text', HEADING_CHARS))
        b['align'] = _align(raw)
    elif kind in ('paragraph', 'quote'):
        b.update(_runs_or_text(raw, 'text', PARAGRAPH_CHARS))
        b['align'] = _align(raw)
    elif kind in ('bullets', 'numbered'):
        b['items'] = [_item(i, n) for n, i in
                      enumerate(items(raw.get('items'), 'items', LIST_ITEMS, required=True), 1)]
    elif kind == 'table':
        columns = [text(c, 'column header', CELL_CHARS, required=True)
                   for c in items(raw.get('columns'), 'columns', TABLE_COLS, required=True)]
        rows = []
        for r, row in enumerate(items(raw.get('rows'), 'rows', TABLE_ROWS, required=True), 1):
            if not isinstance(row, list) or len(row) > len(columns):
                raise SpecError(f'row {r} must be a list of at most {len(columns)} values.')
            rows.append([text(v, f'row {r} cell', CELL_CHARS) for v in row]
                        + [''] * (len(columns) - len(row)))
        b.update(columns=columns, rows=rows,
                 caption=text(raw.get('caption'), 'caption', CAPTION_CHARS))
    elif kind == 'image':
        b['path'] = text(raw.get('path'), 'path', 500, required=True)
        b['caption'] = text(raw.get('caption'), 'caption', CAPTION_CHARS)
    elif kind == 'chart':
        try:
            b['chart'] = build_spec(raw.get('chart'))
        except ChartError as exc:
            raise SpecError(f'chart: {exc}') from None
    return b


def image_paths(spec: dict) -> list[str]:
    return [b['path'] for b in spec['blocks'] if b['type'] == 'image']


def chart_count(spec: dict) -> int:
    return sum(1 for b in spec['blocks'] if b['type'] == 'chart')


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def render(spec: dict, images: dict[str, bytes]) -> bytes:
    theme = THEMES[spec['theme']]
    doc = new_docx()
    _styles(doc, theme)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)

    title = doc.add_paragraph(style='Title')
    _inline(title, spec['title'])
    if spec['subtitle']:
        sub = doc.add_paragraph(style='Subtitle')
        _inline(sub, spec['subtitle'])

    for b in spec['blocks']:
        _BLOCKS[b['type']](doc, b, theme, images)

    doc.core_properties.title = spec['title']
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _styles(doc, theme: Theme) -> None:
    styles = doc.styles
    normal = styles['Normal'].font
    normal.name, normal.size = theme.body_font, Pt(11)
    normal.color.rgb = _rgb(theme.text)
    styles['Normal'].paragraph_format.space_after = Pt(8)
    styles['Normal'].paragraph_format.line_spacing = 1.15

    for name, size, color in (('Title', 28, theme.text), ('Subtitle', 14, theme.muted),
                              ('Heading 1', 18, theme.accent), ('Heading 2', 14, theme.text),
                              ('Heading 3', 12, theme.text)):
        font = styles[name].font
        font.name, font.size = theme.heading_font, Pt(size)
        font.color.rgb = _rgb(color)
        font.bold = name != 'Subtitle'
        font.italic = False
        # The template's heading fonts are theme references (`asciiTheme`),
        # which win over `font.name`; clearing them is what makes it stick.
        rfonts = styles[name].element.rPr.find(qn('w:rFonts'))
        if rfonts is not None:
            for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                rfonts.attrib.pop(qn(attr), None)
    styles['Title'].paragraph_format.space_after = Pt(4)


def _inline(paragraph, body: str) -> None:
    """`body` into `paragraph`, with `**bold**` and `*italic*` spans."""
    for part in _INLINE.split(body):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _preserve(shaped) -> None:
    """Keep edge spaces: without `xml:space="preserve"` Word trims every run
    that starts or ends in a space, gluing runs together."""
    from docx.oxml.ns import qn

    shaped._r.set(qn('xml:space'), 'preserve')


def _inline_runs(paragraph, runs: list) -> None:
    """Typed runs into `paragraph` — what TipTap and the importer produce."""
    for run in runs:
        if run.get('link'):
            _add_link(paragraph, run['link'], run)
        else:
            shaped = paragraph.add_run(run.get('text', ''))
            _style_run(shaped, run)
            _preserve(shaped)


def _style_run(shaped, run: dict) -> None:
    shaped.bold = bool(run.get('bold'))
    shaped.italic = bool(run.get('italic'))
    shaped.underline = bool(run.get('underline'))
    shaped.font.strike = bool(run.get('strike'))
    if run.get('code'):
        shaped.font.name = 'Consolas'


def _add_link(paragraph, url: str, run: dict) -> None:
    """A hyperlink run: python-docx has no paragraph.add_hyperlink, so the
    relationship and element are built by hand (the documented recipe)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), r_id)
    fresh = OxmlElement('w:r')
    element.append(fresh)
    paragraph._p.append(element)
    from docx.text.run import Run

    shaped = Run(fresh, paragraph)
    _style_run(shaped, run)
    shaped.text = run.get('text', '')
    _preserve(shaped)
    shaped.font.color.rgb = _rgb('0563C1')


def _aligned(paragraph, align: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                           'center': WD_ALIGN_PARAGRAPH.CENTER,
                           'right': WD_ALIGN_PARAGRAPH.RIGHT,
                           'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]


def _render_text(paragraph, block) -> None:
    """A block's words, whether stored as runs or as marker text."""
    if block.get('runs') is not None:
        _inline_runs(paragraph, block['runs'])
    else:
        _inline(paragraph, block.get('text', ''))
    if block.get('align') and block['align'] != 'left':
        _aligned(paragraph, block['align'])


def _render_item(paragraph, item: Any) -> None:
    if isinstance(item, dict):
        _render_text(paragraph, item)
    else:
        _inline(paragraph, item)


def _heading(doc, b, theme, images):
    _render_text(doc.add_paragraph(style=f'Heading {b["level"]}'), b)


def _paragraph(doc, b, theme, images):
    _render_text(doc.add_paragraph(), b)


def _list(doc, b, theme, images):
    style = 'List Bullet' if b['type'] == 'bullets' else 'List Number'
    for item in b['items']:
        _render_item(doc.add_paragraph(style=style), item)


def _quote(doc, b, theme, images):
    p = doc.add_paragraph(style='Quote')
    _render_text(p, b)


def _page_break(doc, b, theme, images):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def _write_table(doc, columns: list[str], rows: list[list[str]], theme: Theme,
                 caption: str) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(columns))
    table.style = 'Table Grid'
    for c, header in enumerate(columns):
        cell = table.cell(0, c)
        _shade(cell, theme.accent)
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = _rgb(theme.on_accent)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            if r % 2 == 0:
                _shade(cell, theme.surface)
            _inline(cell.paragraphs[0], value)
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = _rgb(theme.muted)


def _table(doc, b, theme, images):
    _write_table(doc, b['columns'], b['rows'], theme, b['caption'])


def _chart(doc, b, theme, images):
    spec = b['chart']
    series = spec['series']
    xs: list[str] = []
    for s in series:
        for p in s['points']:
            if p['x'] not in xs:
                xs.append(p['x'])
    columns = [spec.get('x_label') or 'Category'] + [s['name'] for s in series]
    lookup = [{p['x']: p['y'] for p in s['points']} for s in series]
    rows = [[x] + [_fmt(m.get(x)) for m in lookup] for x in xs]
    title = doc.add_paragraph()
    run = title.add_run(spec['title'])
    run.bold = True
    _write_table(doc, columns, rows, theme,
                 (spec.get('note') + ' ' if spec.get('note') else '')
                 + '(Chart data shown as a table.)')


def _fmt(value) -> str:
    if value is None:
        return '—'
    return f'{value:,.0f}' if float(value).is_integer() else f'{value:,.2f}'


def _image(doc, b, theme, images):
    data = images.get(b['path'])
    if data:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as img:
            w, h = img.size
        width = Inches(6.5) if w >= h else Inches(4.5)
        doc.add_picture(io.BytesIO(data), width=width)
    if b['caption']:
        p = doc.add_paragraph()
        run = p.add_run(b['caption'])
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = _rgb(theme.muted)


_BLOCKS = {
    'heading': _heading,
    'paragraph': _paragraph,
    'bullets': _list,
    'numbered': _list,
    'table': _table,
    'quote': _quote,
    'image': _image,
    'chart': _chart,
    'page_break': _page_break,
}


# ---------------------------------------------------------------------------
# What is stored besides the bytes
# ---------------------------------------------------------------------------

def preview(spec: dict) -> dict:
    t = THEMES[spec['theme']]
    return {'kind': 'document', 'accent': t.accent, **spec}


def _runs_markdown(runs: list) -> str:
    """Runs back to marker text: bold/italic keep their markers, links keep
    their target, and the rest travels as plain words."""
    parts = []
    for run in runs:
        body = run.get('text', '')
        if run.get('bold'):
            body = f'**{body}**'
        if run.get('italic'):
            body = f'*{body}*'
        if run.get('link'):
            body = f'[{body}]({run["link"]})'
        parts.append(body)
    return ''.join(parts)


def _text_of(block: dict) -> str:
    if block.get('runs') is not None:
        return _runs_markdown(block['runs'])
    return block.get('text', '')


def _item_text(item: Any) -> str:
    return _text_of(item) if isinstance(item, dict) else str(item)


def extract_text(spec: dict) -> str:
    out = [spec['title']]
    if spec['subtitle']:
        out.append(spec['subtitle'])
    out.append('')
    for b in spec['blocks']:
        kind = b['type']
        if kind == 'heading':
            out.append('#' * b['level'] + ' ' + _text_of(b))
        elif kind in ('paragraph', 'quote'):
            out.append(('> ' if kind == 'quote' else '') + _text_of(b))
        elif kind in ('bullets', 'numbered'):
            out.extend(f'{"-" if kind == "bullets" else f"{i}."} {_item_text(item)}'
                       for i, item in enumerate(b['items'], 1))
        elif kind == 'table':
            out.append(' | '.join(b['columns']))
            out.extend(' | '.join(r) for r in b['rows'])
        elif kind == 'chart':
            out.append(f'Chart: {b["chart"]["title"]}')
        elif kind == 'image' and b['caption']:
            out.append(f'[Image: {b["caption"]}]')
        out.append('')
    return '\n'.join(out)
