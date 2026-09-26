"""
Office files as tools: `render_deck`, `render_workbook`, `render_document`.

Each takes a **spec** — slides, sheets, blocks — and our code produces the
file. The model never writes python-pptx code and never lays anything out,
for the reason `render_chart` exists: a model asked to be a layout engine
produces a different-looking file every time, and some of them broken. One
renderer and three themes are one design system.

Three properties are shared, and they are why these tools can run without
asking in chat (unlike `write_file`):

* **They only create.** A name that is taken becomes `name (2).ext`;
  replacing needs `overwrite: true`, and even then the file is replaced in
  place (same id) after keeping what it held as a version
  (`vfs.write_binary`, `inference/versions.py`). So the effect is
  `reversible` in the autonomy ladder's sense — nothing done here is beyond
  the user's own undo.
* **They go through the caller's `FileScope`**, like every other file tool:
  `requires="files"` withholds them where there is no scope, images are read
  through the same walk, and the write is checked against the same prefix.
* **Every limit refuses rather than truncates** (`office/spec.py`), and says
  what to change.

The rendering is CPU work on pure-Python libraries, done in the web process:
bounded by those same limits, and far cheaper than a LibreOffice sidecar on a
box this size.
"""
from __future__ import annotations

import json
import logging
from typing import Any, Dict

from asgiref.sync import sync_to_async

from .charts import render_chart as _register_chart  # noqa: F401 — schema reused below
from .registry import get as _registered
from .registry import tool
from office import deck, diagram, document, edit, pdf, workbook
from office.spec import SpecError
from office.themes import THEME_NAMES

logger = logging.getLogger(__name__)

#: A chart on a slide or in a document takes exactly `render_chart`'s shape.
_CHART_SCHEMA = {
    **_registered('render_chart').schema['function']['parameters'],
    'description': 'A chart in the same shape render_chart takes.',
}

_PATH = {
    'type': 'string',
    'description': (
        'Where to save it, e.g. "/Chat/q3-review{ext}". A bare file name is '
        'saved in your own folder.'
    ),
}
_OVERWRITE = {
    'type': 'boolean',
    'description': (
        'Replace a file already at this path (what it held becomes a version '
        'in its history). Without this, a taken name is saved as "name (2)".'
    ),
}


def _path_schema(ext: str) -> dict:
    return {**_PATH, 'description': _PATH['description'].format(ext=ext)}


# ---------------------------------------------------------------------------
# Shared plumbing
# ---------------------------------------------------------------------------

def _target(scope, raw: Any, ext: str, fallback: str) -> str:
    """The path to save to: the right extension, and in a writable folder.

    A bare name goes in the scope's own write folder (`/Chat/` in chat, the
    agent's home for an agent) — the only place most scopes can write, so
    defaulting there turns "deck.pptx" into a file rather than a refusal.
    """
    from inference.vfs import safe_name

    path = str(raw or '').strip() or f'{safe_name(fallback)[:80] or "untitled"}.{ext}'
    leaf = path.replace('\\', '/').rstrip('/').rsplit('/', 1)[-1]
    if '.' in leaf:
        if leaf.rsplit('.', 1)[-1].lower() != ext:
            raise SpecError(f'This tool writes .{ext} files; the path must end in .{ext}.')
    else:
        path = f'{path.rstrip("/")}.{ext}'
    bare = '/' not in path.strip('/')
    if bare and scope.write_prefix:
        path = '/' + '/'.join(scope.write_prefix) + '/' + path.strip('/')
    return path


def _load_images(scope, paths: list[str]) -> dict[str, bytes]:
    from inference.vfs import read_image

    return {p: read_image(scope, p)[0] for p in dict.fromkeys(paths)}


async def _save(context: Dict, run) -> str:
    """Run `run(scope)` — an async function doing its own thread placement —
    and render its outcome for the model.

    The rule (G5): ORM stays on the run's thread (plain `sync_to_async`),
    CPU and HTTP go to the default pool (`thread_sensitive=False`). `run`
    places each stage itself, so a render never sits on the run's only thread
    blocking that run's ORM calls. Errors map exactly as before.
    """
    from inference.vfs import VfsError

    scope = context.get('file_scope')
    if scope is None:
        return json.dumps({
            'error': 'There is no file workspace here, so nothing can be saved. '
                     'An agent needs file access turned on in its settings.'
        })
    try:
        result = await run(scope)
    except (SpecError, VfsError) as exc:
        return json.dumps({'error': str(exc)})
    except Exception:
        logger.exception('[Office] render failed')
        return json.dumps({'error': 'The file could not be rendered. Simplify it and try again.'})

    if result.get('renamed'):
        result['note'] = (
            f'A file with that name already existed, so this was saved as '
            f'{result["path"]}. Pass overwrite=true to replace the old one instead.'
        )
    result['rendered'] = (
        f'Saved {result["path"]}. The user sees it as a file card they can '
        f'preview and download — do not paste its contents back; say in a '
        f'sentence or two what is in it.'
    )
    return json.dumps(result, default=str)


# ---------------------------------------------------------------------------
# render_workbook
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'render_workbook',
        'description': (
            'Create an Excel workbook (.xlsx) from sheets of typed columns and '
            'rows; the app does all formatting. Use it whenever the user wants '
            'a spreadsheet, a table they will work in, or numbers they will '
            'reuse. Keep calculations live: a value starting with "=" is a '
            'formula (row 1 is the header, data starts on row 2, and {r} means '
            '"this row", e.g. "=B{r}*C{r}"), and totals: true adds a SUM row. '
            'Do not type in totals you computed. For more than a few thousand '
            'rows, compute with execute_python first and save a summary.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': _path_schema('.xlsx'),
                'sheets': {
                    'type': 'array',
                    'description': f'At most {workbook.MAX_SHEETS} sheets.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'name': {'type': 'string', 'description': 'Sheet tab name, at most 31 characters.'},
                            'columns': {
                                'type': 'array',
                                'items': {
                                    'type': 'object',
                                    'properties': {
                                        'header': {'type': 'string'},
                                        'type': {
                                            'type': 'string',
                                            'enum': list(workbook.COLUMN_TYPES),
                                            'description': 'percent takes fractions (0.25 is 25%); date takes YYYY-MM-DD.',
                                        },
                                        'currency': {
                                            'type': 'string',
                                            'enum': list(workbook.CURRENCY_SYMBOLS),
                                            'description': 'For currency columns.',
                                        },
                                    },
                                    'required': ['header'],
                                    'additionalProperties': False,
                                },
                            },
                            'rows': {
                                'type': 'array',
                                'description': f'One list of values per row, in column order. At most {workbook.MAX_ROWS}.',
                                'items': {'type': 'array', 'items': {'type': ['string', 'number', 'boolean', 'null']}},
                            },
                            'totals': {
                                'type': 'boolean',
                                'description': 'Add a Total row summing every number/integer/currency column.',
                            },
                            'chart': {
                                'type': 'object',
                                'description': 'A native Excel chart over this sheet\'s own columns.',
                                'properties': {
                                    'kind': {'type': 'string', 'enum': ['bar', 'column', 'line', 'area', 'scatter', 'pie']},
                                    'title': {'type': 'string'},
                                    'x': {'type': 'string', 'description': 'Header of the category column.'},
                                    'y': {'type': 'array', 'items': {'type': 'string'}, 'description': 'Headers of the numeric columns to plot.'},
                                    'x_label': {'type': 'string'},
                                    'y_label': {'type': 'string'},
                                    'stacked': {'type': 'boolean'},
                                },
                                'required': ['kind', 'title', 'x', 'y'],
                                'additionalProperties': False,
                            },
                        },
                        'required': ['name', 'columns', 'rows'],
                        'additionalProperties': False,
                    },
                },
                'overwrite': _OVERWRITE,
            },
            'required': ['sheets'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='reversible')
async def render_workbook(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        spec = workbook.validate(
            args,
            max_sheets=await alimit(context, 'render_workbook', 'maxSheets'),
            max_rows=await alimit(context, 'render_workbook', 'maxRows'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import write_binary

        path = _target(scope, args.get('path'), 'xlsx', spec['sheets'][0]['name'])
        data, warnings = await sync_to_async(workbook.render, thread_sensitive=False)(spec)
        result = await sync_to_async(write_binary)(
            scope, path, data, text=workbook.extract_text(spec),
            spec=workbook.preview(spec), overwrite=bool(args.get('overwrite')))
        result['sheets'] = [{'name': s['name'], 'rows': len(s['rows'])} for s in spec['sheets']]
        if warnings:
            result['warnings'] = warnings
        return result

    return await _save(context, run)


# ---------------------------------------------------------------------------
# render_deck
# ---------------------------------------------------------------------------

_BULLETS = {
    'type': 'array',
    'items': {'type': 'string'},
    'description': (
        f'At most {deck.MAX_BULLETS}, each under {deck.BULLET_CHARS} characters. '
        'Start one with "- " to make it a sub-point. **bold** works.'
    ),
}


@tool({
    'type': 'function',
    'function': {
        'name': 'render_deck',
        'description': (
            'Create a PowerPoint deck (.pptx). You choose what each slide says '
            'and its layout; the app does all design, so never describe fonts, '
            'colours or positions. Layouts: title, section, bullets, '
            'two_column, chart (a native, editable chart), image (a picture '
            'from the user\'s files), table, quote, stats (up to 4 big '
            'numbers), closing. One idea per slide; a slide that will not fit '
            'is refused — split it rather than cramming. Put what the presenter '
            'should say in notes.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': _path_schema('.pptx'),
                'title': {'type': 'string', 'description': 'The deck\'s title.'},
                'theme': {'type': 'string', 'enum': list(THEME_NAMES),
                          'description': 'clean (default), bold, or dark.'},
                'slides': {
                    'type': 'array',
                    'description': f'At most {deck.MAX_SLIDES}.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'layout': {'type': 'string', 'enum': list(deck.LAYOUTS)},
                            'title': {'type': 'string', 'description': f'Under {deck.TITLE_CHARS} characters. Every layout but quote.'},
                            'subtitle': {'type': 'string', 'description': 'title, section, closing.'},
                            'bullets': _BULLETS,
                            'left': {
                                'type': 'object', 'description': 'two_column.',
                                'properties': {'heading': {'type': 'string'}, 'bullets': _BULLETS},
                                'required': ['bullets'], 'additionalProperties': False,
                            },
                            'right': {
                                'type': 'object', 'description': 'two_column.',
                                'properties': {'heading': {'type': 'string'}, 'bullets': _BULLETS},
                                'required': ['bullets'], 'additionalProperties': False,
                            },
                            'chart': _CHART_SCHEMA,
                            'image': {'type': 'string', 'description': 'image: path of an image in the user\'s files.'},
                            'caption': {'type': 'string', 'description': 'chart, image: one line under it.'},
                            'columns': {'type': 'array', 'items': {'type': 'string'},
                                        'description': f'table: at most {deck.TABLE_COLS}.'},
                            'rows': {'type': 'array', 'items': {'type': 'array', 'items': {'type': 'string'}},
                                     'description': f'table: at most {deck.TABLE_ROWS}.'},
                            'quote': {'type': 'string'},
                            'attribution': {'type': 'string'},
                            'stats': {
                                'type': 'array', 'description': 'stats: 1 to 4.',
                                'items': {
                                    'type': 'object',
                                    'properties': {
                                        'value': {'type': 'string', 'description': 'e.g. "₹4.2 Cr", "38%".'},
                                        'label': {'type': 'string'},
                                    },
                                    'required': ['value', 'label'],
                                    'additionalProperties': False,
                                },
                            },
                            'notes': {'type': 'string', 'description': 'Speaker notes.'},
                        },
                        'required': ['layout'],
                        'additionalProperties': False,
                    },
                },
                'overwrite': _OVERWRITE,
            },
            'required': ['slides'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='reversible')
async def render_deck(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        spec = deck.validate(
            args, max_slides=await alimit(context, 'render_deck', 'maxSlides'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import write_binary

        path = _target(scope, args.get('path'), 'pptx', spec['title'] or 'deck')
        images = await sync_to_async(_load_images)(scope, deck.image_paths(spec))
        data = await sync_to_async(deck.render, thread_sensitive=False)(spec, images)
        result = await sync_to_async(write_binary)(
            scope, path, data, text=deck.extract_text(spec),
            spec=deck.preview(spec), overwrite=bool(args.get('overwrite')))
        result['slides'] = len(spec['slides'])
        return result

    return await _save(context, run)


# ---------------------------------------------------------------------------
# render_document
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'render_document',
        'description': (
            'Create a Word document (.docx) from a title and a list of blocks: '
            'heading, paragraph, bullets, numbered, table, quote, image, chart, '
            'page_break. The app does all styling. Use it when the user wants '
            'a report, memo or letter as a Word file; for notes they will only '
            'read here, write markdown with write_file instead. **bold** and '
            '*italic* work inside text. A chart block is saved as its data '
            'table (Word files get no drawn charts).'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': _path_schema('.docx'),
                'title': {'type': 'string'},
                'subtitle': {'type': 'string'},
                'theme': {'type': 'string', 'enum': list(document.DOC_THEMES)},
                'blocks': {
                    'type': 'array',
                    'description': f'At most {document.MAX_BLOCKS}.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'type': {'type': 'string', 'enum': list(document.BLOCK_TYPES)},
                            'text': {'type': 'string', 'description': 'heading, paragraph, quote.'},
                            'level': {'type': 'integer', 'enum': [1, 2, 3], 'description': 'heading.'},
                            'items': {'type': 'array', 'items': {'type': 'string'},
                                      'description': 'bullets, numbered.'},
                            'columns': {'type': 'array', 'items': {'type': 'string'}, 'description': 'table.'},
                            'rows': {'type': 'array', 'items': {'type': 'array', 'items': {'type': 'string'}},
                                     'description': 'table.'},
                            'caption': {'type': 'string', 'description': 'table, image.'},
                            'path': {'type': 'string', 'description': 'image: path of an image in the user\'s files.'},
                            'chart': _CHART_SCHEMA,
                        },
                        'required': ['type'],
                        'additionalProperties': False,
                    },
                },
                'overwrite': _OVERWRITE,
            },
            'required': ['title', 'blocks'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='reversible')
async def render_document(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        spec = document.validate(
            args, max_blocks=await alimit(context, 'render_document', 'maxBlocks'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import write_binary

        path = _target(scope, args.get('path'), 'docx', spec['title'])
        images = await sync_to_async(_load_images)(scope, document.image_paths(spec))
        data = await sync_to_async(document.render, thread_sensitive=False)(spec, images)
        result = await sync_to_async(write_binary)(
            scope, path, data, text=document.extract_text(spec),
            spec=document.preview(spec), overwrite=bool(args.get('overwrite')))
        result['blocks'] = len(spec['blocks'])
        if document.chart_count(spec):
            result['charts_as_tables'] = document.chart_count(spec)
        return result

    return await _save(context, run)


# ---------------------------------------------------------------------------
# render_pdf — the same blocks as render_document, as the format people send
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'render_pdf',
        'description': (
            'Create a PDF from the same blocks render_document takes — heading, '
            'paragraph, bullets, numbered, table, quote, image, chart, page_break. '
            'Use it when the file is to be sent, printed or attached rather than '
            'edited; use render_document when the user will edit it in Word. A '
            'chart block is drawn as its data table, as in Word.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': _path_schema('.pdf'),
                'title': {'type': 'string'},
                'subtitle': {'type': 'string'},
                'theme': {'type': 'string', 'enum': list(document.DOC_THEMES)},
                'blocks': {
                    'type': 'array',
                    'description': f'At most {document.MAX_BLOCKS}. Same shape as render_document.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'type': {'type': 'string', 'enum': list(document.BLOCK_TYPES)},
                            'text': {'type': 'string', 'description': 'heading, paragraph, quote.'},
                            'level': {'type': 'integer', 'enum': [1, 2, 3], 'description': 'heading.'},
                            'items': {'type': 'array', 'items': {'type': 'string'},
                                      'description': 'bullets, numbered.'},
                            'columns': {'type': 'array', 'items': {'type': 'string'}, 'description': 'table.'},
                            'rows': {'type': 'array', 'items': {'type': 'array', 'items': {'type': 'string'}},
                                     'description': 'table.'},
                            'caption': {'type': 'string', 'description': 'table, image.'},
                            'path': {'type': 'string', 'description': "image: a path in the user's files."},
                            'chart': _CHART_SCHEMA,
                        },
                        'required': ['type'],
                        'additionalProperties': False,
                    },
                },
                'overwrite': _OVERWRITE,
            },
            'required': ['title', 'blocks'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='reversible')
async def render_pdf(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        spec = document.validate(
            args, max_blocks=await alimit(context, 'render_pdf', 'maxBlocks'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import write_binary

        target = _target(scope, args.get('path'), 'pdf', spec['title'])
        images = await sync_to_async(_load_images)(scope, document.image_paths(spec))
        data = await sync_to_async(pdf.render, thread_sensitive=False)(spec, images)
        result = await sync_to_async(write_binary)(
            scope, target, data, text=document.extract_text(spec),
            spec=document.preview(spec), overwrite=bool(args.get('overwrite')))
        result['blocks'] = len(spec['blocks'])
        if document.chart_count(spec):
            result['charts_as_tables'] = document.chart_count(spec)
        return result

    return await _save(context, run)


# ---------------------------------------------------------------------------
# edit_workbook — change a workbook without re-emitting it
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'edit_workbook',
        'description': (
            'Change a workbook that already exists, keeping everything else — '
            'other sheets, charts and the formulas you do not touch. Prefer '
            'this over render_workbook for any change to an existing file: '
            're-rendering means re-typing every row you did not mean to '
            'change. Values starting with "=" are formulas. Structural edits '
            'run first, so set_cells coordinates refer to the sheet after any '
            'insert/delete in the same call.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': {'type': 'string', 'description': 'The .xlsx to change.'},
                'sheet': {'type': 'string', 'description': 'Sheet name. Defaults to the first.'},
                'append_rows': {
                    'type': 'array',
                    'description': 'Rows added after the last row that has data, in column order.',
                    'items': {'type': 'array', 'items': {'type': ['string', 'number', 'boolean', 'null']}},
                },
                'set_cells': {
                    'type': 'array',
                    'description': 'Individual cells to overwrite.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'cell': {'type': 'string', 'description': 'A reference like B7.'},
                            'value': {'type': ['string', 'number', 'boolean', 'null']},
                        },
                        'required': ['cell', 'value'],
                        'additionalProperties': False,
                    },
                },
                'insert_rows': {
                    'type': 'object',
                    'description': 'Insert blank rows mid-sheet; formulas pointing past them shift along.',
                    'properties': {
                        'row': {'type': 'integer', 'description': '1-based row to insert at.'},
                        'count': {'type': 'integer', 'description': 'How many (default 1).'},
                    },
                    'required': ['row'],
                    'additionalProperties': False,
                },
                'delete_rows': {
                    'type': 'object',
                    'description': 'Delete rows mid-sheet; references into them become #REF! loudly.',
                    'properties': {
                        'row': {'type': 'integer', 'description': '1-based first row to delete.'},
                        'count': {'type': 'integer', 'description': 'How many (default 1).'},
                    },
                    'required': ['row'],
                    'additionalProperties': False,
                },
                'insert_cols': {
                    'type': 'object',
                    'description': 'Insert blank columns mid-sheet; formulas shift along.',
                    'properties': {
                        'col': {'type': 'string', 'description': 'Column like C (or 3).'},
                        'count': {'type': 'integer', 'description': 'How many (default 1).'},
                    },
                    'required': ['col'],
                    'additionalProperties': False,
                },
                'delete_cols': {
                    'type': 'object',
                    'description': 'Delete columns mid-sheet.',
                    'properties': {
                        'col': {'type': 'string', 'description': 'Column like C (or 3).'},
                        'count': {'type': 'integer', 'description': 'How many (default 1).'},
                    },
                    'required': ['col'],
                    'additionalProperties': False,
                },
                'format': {
                    'type': 'object',
                    'description': 'Paint one range.',
                    'properties': {
                        'range': {'type': 'string', 'description': 'A range like A1:B2.'},
                        'style': {
                            'type': 'object',
                            'description': 'bold, italic, wrap (true/false); font_size, font_name, '
                                           'font_color (#rrggbb), fill (#rrggbb), number_format; '
                                           'alignment (left/center/right/justify/top/bottom/middle); '
                                           'border {top/right/bottom/left: thin/medium/…}.',
                        },
                    },
                    'required': ['range', 'style'],
                    'additionalProperties': False,
                },
                'freeze': {
                    'type': ['string', 'null'],
                    'description': 'Freeze rows above and columns left of a cell like B2 (null unfreezes).',
                },
                'widths': {
                    'type': 'object',
                    'description': 'Column widths, e.g. {"A": 20}.',
                    'additionalProperties': {'type': 'number'},
                },
            },
            'required': ['path'],
            'additionalProperties': False,
        },
    },
}, requires='files', sensitive=True, effect='reversible')
async def edit_workbook(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        change = edit.validate(
            args, max_edits=await alimit(context, 'edit_workbook', 'maxEdits'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import VfsError, read_binary, write_binary

        path = str(args.get('path') or '').strip()
        if not path.lower().endswith('.xlsx'):
            raise SpecError('edit_workbook changes .xlsx files; give the path of one.')
        try:
            data = await sync_to_async(read_binary)(scope, path)
        except VfsError:
            raise
        updated, report = await sync_to_async(edit.apply, thread_sensitive=False)(
            data, change)
        text = await sync_to_async(_workbook_text, thread_sensitive=False)(updated)
        result = await sync_to_async(write_binary)(
            scope, path, updated, text=text, overwrite=True)
        result.update(report)
        return result

    return await _save(context, run)


def _workbook_text(data: bytes) -> str:
    """The edited file's searchable text, read back from what was written."""
    import io

    from inference.utils import extract_xlsx_text

    return extract_xlsx_text(io.BytesIO(data))


# ---------------------------------------------------------------------------
# read_workbook — values plus formulas, for reading before editing
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'read_workbook',
        'description': (
            'Read a sheet of a workbook that already exists: every cell as its '
            'calculated value, with the formula beside it wherever one is '
            'stored. Read before edit_workbook rather than guessing what is '
            'in the file. Values starting with "=" in the formulas are live '
            'formulas; a value of null with no formula is an empty cell.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': {'type': 'string', 'description': 'The .xlsx to read.'},
                'sheet': {'type': 'string', 'description': 'Sheet name. Defaults to the first.'},
                'range': {'type': 'string', 'description': 'A range like A1:D20. Defaults to the used area.'},
            },
            'required': ['path'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='read', parallel=True)
async def read_workbook(args: Dict, context: Dict) -> str:
    async def run(scope):
        from inference.vfs import VfsError, read_binary

        path = str(args.get('path') or '').strip()
        if not path.lower().endswith('.xlsx'):
            return json.dumps({'error': 'read_workbook reads .xlsx files; give the path of one.'})
        try:
            data = await sync_to_async(read_binary)(scope, path)
        except VfsError as exc:
            return json.dumps({'error': str(exc)})
        try:
            result = await sync_to_async(_read_sheet, thread_sensitive=False)(
                data, args.get('sheet'), args.get('range'))
        except SpecError as exc:
            return json.dumps({'error': str(exc)})
        return json.dumps(result, default=str)

    scope = context.get('file_scope')
    if scope is None:
        return json.dumps({'error': 'There is no file workspace here.'})
    try:
        return await run(scope)
    except Exception:
        logger.exception('[Office] read_workbook failed')
        return json.dumps({'error': 'The workbook could not be read.'})


def _read_sheet(data: bytes, sheet: Any, cell_range: Any) -> dict:
    """One sheet's cells as values plus formulas, optionally clipped to a range."""
    import re

    from office import formulas

    wb = formulas.workbook(data)
    try:
        name = str(sheet or '').strip() or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise SpecError(f'No sheet called {name!r}. This workbook has: {", ".join(wb.sheetnames)}.')
        ws = wb[name]
        if cell_range:
            bounds = str(cell_range).strip().upper()
            if not re.fullmatch(r'\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}'
                                r'(:\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?', bounds):
                raise SpecError(f'range {bounds!r} is not like A1:D20.')
            # Clipped to the used area: `ws[...]` creates a cell object per
            # coordinate, and A1:XFD9999999 would create billions of them.
            from openpyxl.utils.cell import range_boundaries

            c1, r1, c2, r2 = range_boundaries(bounds.replace('$', ''))
            cells = ws.iter_rows(min_row=r1, max_row=min(r2 or r1, ws.max_row or 0),
                                 min_col=c1, max_col=min(c2 or c1, ws.max_column or 0))
        else:
            cells = ws.iter_rows(min_row=1, max_row=ws.max_row or 0,
                                 max_col=ws.max_column or 0)
        rows: list[list] = []
        formula_rows: list[list] = []
        for row in cells:
            cells_in_row = row if isinstance(row, tuple) else (row,)
            values: list = []
            formulae: list = []
            for cell in cells_in_row:
                raw = cell.value
                if isinstance(raw, str) and raw.startswith('='):
                    try:
                        values.append(formulas.cell(wb, name, cell.coordinate))
                    except formulas.FormulaError:
                        values.append(None)
                    formulae.append(raw)
                else:
                    values.append(raw)
                    formulae.append(None)
            while values and values[-1] is None and formulae[-1] is None:
                values.pop()
                formulae.pop()
            rows.append([formulas.json_value(v) for v in values])
            formula_rows.append(formulae)
        while rows and not any(rows[-1]):
            rows.pop()
            formula_rows.pop()
        return {'sheet': name, 'sheets': wb.sheetnames,
                'values': rows, 'formulas': formula_rows}
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# render_diagram — boxes and arrows as data
# ---------------------------------------------------------------------------

@tool({
    'type': 'function',
    'function': {
        'name': 'render_diagram',
        'description': (
            'Draw a flow or architecture diagram from nodes and edges and save it '
            'as an SVG in the user\'s files. The app does the layout (left to '
            'right, one column per step), so describe what connects to what and '
            'never write SVG or Mermaid yourself. Good for pipelines, processes '
            'and system maps; use render_chart for numbers. The SVG opens in a '
            'browser and goes into documents — it cannot be put on a slide yet.'
        ),
        'parameters': {
            'type': 'object',
            'properties': {
                'path': _path_schema('.svg'),
                'title': {'type': 'string', 'description': 'What the diagram shows.'},
                'theme': {'type': 'string', 'enum': list(THEME_NAMES)},
                'nodes': {
                    'type': 'array',
                    'description': f'At most {diagram.MAX_NODES}.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'id': {'type': 'string', 'description': 'Short, unique; edges refer to it.'},
                            'label': {'type': 'string', 'description': 'What is shown in the box.'},
                            'shape': {'type': 'string', 'enum': list(diagram.SHAPES),
                                      'description': 'diamond for a decision.'},
                            'accent': {'type': 'boolean', 'description': 'Highlight this one.'},
                        },
                        'required': ['id'],
                        'additionalProperties': False,
                    },
                },
                'edges': {
                    'type': 'array',
                    'description': f'At most {diagram.MAX_EDGES}.',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'from': {'type': 'string'},
                            'to': {'type': 'string'},
                            'label': {'type': 'string', 'description': 'What the arrow means.'},
                        },
                        'required': ['from', 'to'],
                        'additionalProperties': False,
                    },
                },
                'overwrite': _OVERWRITE,
            },
            'required': ['nodes'],
            'additionalProperties': False,
        },
    },
}, requires='files', effect='reversible')
async def render_diagram(args: Dict, context: Dict) -> str:
    from tools_config.overlay import alimit

    try:
        spec = diagram.validate(
            args,
            max_nodes=await alimit(context, 'render_diagram', 'maxNodes'),
            max_edges=await alimit(context, 'render_diagram', 'maxEdges'))
    except SpecError as exc:
        return json.dumps({'error': str(exc)})

    async def run(scope):
        from inference.vfs import write_binary

        target = _target(scope, args.get('path'), 'svg', spec['title'] or 'diagram')
        data = await sync_to_async(diagram.render, thread_sensitive=False)(spec)
        result = await sync_to_async(write_binary)(
            scope, target, data, text=diagram.extract_text(spec),
            spec={'kind': 'diagram', **spec}, overwrite=bool(args.get('overwrite')))
        result['nodes'] = len(spec['nodes'])
        result['edges'] = len(spec['edges'])
        return result

    return await _save(context, run)


OFFICE_TOOLS = ('render_deck', 'render_workbook', 'render_document', 'render_pdf',
                'edit_workbook', 'read_workbook', 'render_diagram')
