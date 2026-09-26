"""
Version history (inference/versions.py) and export (inference/export.py).

What these pin:
* every overwrite keeps what the file held — app saves, agent writes, and a
  restore — and a burst of saves from one source makes one version;
* an agent overwriting a deck or workbook keeps the **same document id**
  (it used to trash the file and mint a new one, pulling it out from under
  anyone who had it open in an app);
* an agent's `write_file` over an uploaded text file reaches its bytes;
* restore brings the spec back with the bytes, and is itself undoable;
* another user can neither list nor restore someone else's versions;
* export offers exactly the formats it can make, and makes them.
"""
from __future__ import annotations

import io
import shutil
import tempfile
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from inference import versions, vfs
from inference.models import Document, DocumentVersion

User = get_user_model()

_MEDIA = tempfile.mkdtemp(prefix='versions-')


def _stream(resp) -> bytes:
    return b''.join(resp.streaming_content) if getattr(resp, 'streaming', False) else resp.content


@override_settings(MEDIA_ROOT=_MEDIA)
class VersionCase(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(_MEDIA, ignore_errors=True)

    def setUp(self):
        self.owner = User.objects.create_user('owner', 'o@example.com', 'pw')
        self.other = User.objects.create_user('other', 't@example.com', 'pw')
        self.me = APIClient()
        self.me.force_authenticate(self.owner)
        self.them = APIClient()
        self.them.force_authenticate(self.other)
        self.scope = vfs.build_scope(self.owner, 'full')

    def new(self, name, **extra):
        resp = self.me.post('/api/inference/documents/new/', {'name': name, **extra}, format='json')
        self.assertEqual(resp.status_code, 201, resp.content)
        return resp.json()

    def save_text(self, doc_id, content):
        resp = self.me.patch(f'/api/inference/documents/{doc_id}/content/',
                             {'content': content}, format='json')
        self.assertEqual(resp.status_code, 200, resp.content)
        return resp.json()

    def age_versions(self, doc_id, seconds=3600):
        DocumentVersion.objects.filter(document_id=doc_id).update(
            created_at=timezone.now() - timedelta(seconds=seconds))


class SnapshotTests(VersionCase):
    def test_an_app_save_keeps_what_the_file_held(self):
        doc = self.new('notes.md', content='first')
        self.save_text(doc['id'], 'second')
        kept = DocumentVersion.objects.get(document_id=doc['id'])
        self.assertEqual(kept.content_text, 'first')
        self.assertEqual(kept.source, 'app')

    def test_a_burst_of_saves_makes_one_version(self):
        doc = self.new('notes.md', content='v0')
        for n in range(1, 6):
            self.save_text(doc['id'], f'v{n}')
        rows = DocumentVersion.objects.filter(document_id=doc['id'])
        self.assertEqual(rows.count(), 1)
        # The one kept is the state before the burst began.
        self.assertEqual(rows.get().content_text, 'v0')

    def test_a_later_session_makes_another(self):
        doc = self.new('notes.md', content='v0')
        self.save_text(doc['id'], 'v1')
        self.age_versions(doc['id'])
        self.save_text(doc['id'], 'v2')
        self.assertEqual(DocumentVersion.objects.filter(document_id=doc['id']).count(), 2)

    def test_an_agent_write_after_an_app_save_gets_its_own_version(self):
        doc = self.new('notes.md', content='mine')
        self.save_text(doc['id'], 'mine, edited')
        vfs.write_file(self.scope, '/notes.md', 'the agent rewrote it')
        sources = list(DocumentVersion.objects.filter(document_id=doc['id'])
                       .values_list('source', flat=True))
        self.assertEqual(sorted(sources), ['agent', 'app'])

    def test_an_app_save_after_an_agent_write_keeps_the_agents_text(self):
        # Coalescing by source reached past the agent's write: the next app
        # save kept no copy of it.
        doc = self.new('notes.md', content='mine')
        self.save_text(doc['id'], 'mine, edited')
        vfs.write_file(self.scope, '/notes.md', 'the agent rewrote it')
        self.save_text(doc['id'], 'mine again')
        texts = list(DocumentVersion.objects.filter(document_id=doc['id'])
                     .values_list('content_text', flat=True))
        self.assertIn('the agent rewrote it', texts)

    def test_the_cap_prunes_the_oldest_with_its_blob(self):
        from workflow_backend import thresholds

        doc = self.new('notes.md', content='v0')
        with self.settings():
            old = thresholds.FILE_VERSIONS_KEPT
            thresholds.FILE_VERSIONS_KEPT = 3
            try:
                for n in range(1, 7):
                    self.age_versions(doc['id'])
                    self.save_text(doc['id'], f'v{n}')
            finally:
                thresholds.FILE_VERSIONS_KEPT = old
        self.assertEqual(DocumentVersion.objects.filter(document_id=doc['id']).count(), 3)

    def test_a_failure_to_keep_a_version_never_fails_the_save(self):
        doc = self.new('notes.md', content='v0')
        from unittest import mock

        with mock.patch('inference.versions.DocumentVersion.save', side_effect=RuntimeError('disk full')):
            self.save_text(doc['id'], 'v1')
        self.assertEqual(Document.objects.get(id=doc['id']).content_text, 'v1')


class AgentOverwriteTests(VersionCase):
    def test_write_file_over_an_uploaded_file_reaches_its_bytes(self):
        doc = Document(user=self.owner, name='a.md', file_type='md', content_text='old',
                       file_size=3, status='indexed')
        doc.file.save('a.md', ContentFile(b'old'), save=True)
        vfs.write_file(self.scope, '/a.md', 'rewritten by an agent')
        resp = self.me.get(f'/api/inference/documents/{doc.id}/download/')
        self.assertEqual(_stream(resp), b'rewritten by an agent')

    def test_overwriting_a_binary_keeps_its_id_and_a_version(self):
        doc = self.new('Budget.xlsx')
        from office import workbook

        spec = workbook.validate({'sheets': [{'name': 'S', 'columns': ['A'], 'rows': [['x']]}]})
        data, _ = workbook.render(spec)
        result = vfs.write_binary(self.scope, '/Budget.xlsx', data, overwrite=True)
        self.assertEqual(result['document_id'], doc['id'])
        self.assertTrue(result['replaced'])
        self.assertEqual(Document.objects.filter(user=self.owner, name='Budget.xlsx').count(), 1)
        self.assertEqual(DocumentVersion.objects.get(document_id=doc['id']).source, 'agent')

    def test_overwrite_without_a_spec_drops_the_stale_one(self):
        deck = self.new('Pitch.pptx')
        self.assertIn('spec', Document.objects.get(id=deck['id']).metadata)
        vfs.write_binary(self.scope, '/Pitch.pptx', b'PK\x03\x04not really a deck', overwrite=True)
        self.assertNotIn('spec', Document.objects.get(id=deck['id']).metadata)


class RestoreTests(VersionCase):
    def test_restore_puts_the_text_back_and_is_itself_undoable(self):
        doc = self.new('notes.md', content='original')
        self.save_text(doc['id'], 'changed')
        vid = DocumentVersion.objects.get(document_id=doc['id']).id
        resp = self.me.post(f'/api/inference/documents/{doc["id"]}/versions/{vid}/restore/')
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(Document.objects.get(id=doc['id']).content_text, 'original')
        before_restore = DocumentVersion.objects.get(document_id=doc['id'], source='restore')
        self.assertEqual(before_restore.content_text, 'changed')

    def test_restore_keeps_a_pending_autosave_as_a_version(self):
        doc = self.new('Plan.docx')
        etag = doc['updated_at']
        for words in ('first draft', 'latest words'):
            resp = self.me.post(
                f'/api/inference/documents/{doc["id"]}/draft/',
                {'spec': {'title': 'Plan', 'blocks': [{'type': 'paragraph', 'text': words}]},
                 'expected_updated_at': etag}, format='json')
            self.assertEqual(resp.status_code, 200, resp.content)
            etag = resp.json()['updated_at']
            if words == 'first draft':
                from inference import drafts

                self.assertTrue(drafts.maybe_render_draft(doc['id'], force=True))
                etag = Document.objects.get(id=doc['id']).updated_at.isoformat()
        vid = DocumentVersion.objects.filter(document_id=doc['id']).earliest('created_at').id
        resp = self.me.post(f'/api/inference/documents/{doc["id"]}/versions/{vid}/restore/')
        self.assertEqual(resp.status_code, 200, resp.content)
        before_restore = DocumentVersion.objects.get(document_id=doc['id'], source='restore')
        self.assertIn('latest words', before_restore.content_text)

    def test_restore_brings_the_spec_back_with_the_bytes(self):
        deck = self.new('Pitch.pptx')
        original = Document.objects.get(id=deck['id']).metadata['spec']
        slides = [{'layout': 'bullets', 'title': 'New', 'bullets': ['one']}]
        resp = self.me.post(f'/api/inference/documents/{deck["id"]}/office/',
                            {'spec': {'slides': slides}}, format='json')
        self.assertEqual(resp.status_code, 200, resp.content)
        vid = DocumentVersion.objects.get(document_id=deck['id']).id
        self.me.post(f'/api/inference/documents/{deck["id"]}/versions/{vid}/restore/')
        self.assertEqual(Document.objects.get(id=deck['id']).metadata['spec'], original)

    def test_a_stale_restore_is_refused(self):
        doc = self.new('notes.md', content='a')
        self.save_text(doc['id'], 'b')
        vid = DocumentVersion.objects.get(document_id=doc['id']).id
        resp = self.me.post(f'/api/inference/documents/{doc["id"]}/versions/{vid}/restore/',
                            {'expected_updated_at': doc['updated_at']}, format='json')
        self.assertEqual(resp.status_code, 412)

    def test_the_agent_tools_list_and_restore(self):
        doc = self.new('notes.md', content='draft one')
        vfs.write_file(self.scope, '/notes.md', 'draft two')
        listed = vfs.file_versions(self.scope, '/notes.md')
        self.assertEqual(len(listed['versions']), 1)
        vfs.restore_file_version(self.scope, '/notes.md', listed['versions'][0]['version_id'])
        self.assertEqual(Document.objects.get(id=doc['id']).content_text, 'draft one')

    def test_a_read_only_scope_cannot_restore(self):
        self.new('notes.md', content='a')
        vfs.write_file(self.scope, '/notes.md', 'b')
        vid = vfs.file_versions(self.scope, '/notes.md')['versions'][0]['version_id']
        readonly = vfs.build_scope(self.owner, 'readonly')
        with self.assertRaises(vfs.VfsError):
            vfs.restore_file_version(readonly, '/notes.md', vid)


class IsolationTests(VersionCase):
    def test_another_user_sees_nothing(self):
        doc = self.new('notes.md', content='secret v1')
        self.save_text(doc['id'], 'secret v2')
        vid = DocumentVersion.objects.get(document_id=doc['id']).id
        self.assertEqual(self.them.get(f'/api/inference/documents/{doc["id"]}/versions/').status_code, 404)
        self.assertEqual(self.them.get(
            f'/api/inference/documents/{doc["id"]}/versions/{vid}/download/').status_code, 404)
        self.assertEqual(self.them.post(
            f'/api/inference/documents/{doc["id"]}/versions/{vid}/restore/').status_code, 404)

    def test_the_listing_and_download_work_for_the_owner(self):
        doc = self.new('notes.md', content='kept text')
        self.save_text(doc['id'], 'newer')
        listed = self.me.get(f'/api/inference/documents/{doc["id"]}/versions/').json()['versions']
        self.assertEqual(len(listed), 1)
        resp = self.me.get(f'/api/inference/documents/{doc["id"]}/versions/{listed[0]["id"]}/download/')
        self.assertEqual(_stream(resp), b'kept text')


class ExportTests(VersionCase):
    def export(self, doc_id, fmt=None):
        url = f'/api/inference/documents/{doc_id}/export/'
        return self.me.get(url, {'to': fmt} if fmt else {})

    def test_the_formats_offered_are_the_formats_made(self):
        from inference.export import FORMATS

        samples = {'docx': 'Plan.docx', 'md': 'n.md', 'txt': 'n.txt', 'pptx': 'P.pptx',
                   'xlsx': 'B.xlsx', 'csv': 'c.csv'}
        for ftype, name in samples.items():
            doc = self.new(name, content='# Hello\n\n- one\n- two\n' if ftype in ('md', 'txt') else
                           ('a,b\n1,2\n' if ftype == 'csv' else ''))
            offered = self.export(doc['id']).json()['formats']
            self.assertEqual(tuple(offered), FORMATS[ftype])
            for fmt in offered:
                resp = self.export(doc['id'], fmt)
                body = _stream(resp)
                self.assertEqual(resp.status_code, 200, f'{name} -> {fmt}: {body[:300]}')
                self.assertTrue(body, f'{name} -> {fmt} was empty')

    def test_a_pdf_is_a_pdf_and_a_docx_is_a_zip(self):
        md = self.new('notes.md', content='# Title\n\nSome **bold** text.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n')
        self.assertTrue(_stream(self.export(md['id'], 'pdf')).startswith(b'%PDF'))
        self.assertTrue(_stream(self.export(md['id'], 'docx')).startswith(b'PK'))

    def test_an_unoffered_format_is_refused_with_the_offer(self):
        doc = self.new('notes.md', content='x')
        resp = self.export(doc['id'], 'xlsx')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('pdf', resp.json()['error'])

    def test_export_file_saves_beside_and_never_overwrites(self):
        self.new('notes.md', content='# Report\n\nBody.')
        first = vfs.export_file(self.scope, '/notes.md', 'pdf')
        second = vfs.export_file(self.scope, '/notes.md', 'pdf')
        self.assertEqual(first['path'], '/notes.pdf')
        self.assertEqual(second['path'], '/notes (2).pdf')

    def test_a_csv_round_trips_through_a_workbook(self):
        import openpyxl

        doc = self.new('c.csv', content='name,value\nrent,1200\n')
        data = _stream(self.export(doc['id'], 'xlsx'))
        ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
        self.assertEqual([c.value for c in ws[2]], ['rent', 1200])


class MarkdownBlocksTests(TestCase):
    def test_the_reader_keeps_every_kind(self):
        from inference.text_blocks import markdown_blocks

        title, blocks = markdown_blocks(
            '# Doc\n\nIntro _soft_ text.\n\n## Part\n\n- a\n- [x] done\n\n1. first\n\n'
            '> quoted\n\n```\ncode\n```\n\n| h1 | h2 |\n|---|---|\n| x | y |\n')
        self.assertEqual(title, 'Doc')
        kinds = [b['type'] for b in blocks]
        self.assertEqual(kinds, ['paragraph', 'heading', 'bullets', 'numbered', 'quote',
                                 'paragraph', 'table'])
        self.assertIn('*soft*', blocks[0]['text'])
        self.assertEqual(blocks[2]['items'][1], '☑ done')

    def test_long_text_is_split_not_cut(self):
        from inference.text_blocks import PARAGRAPH_SPLIT, text_blocks

        text = ('A sentence here. ' * 1000).strip()
        blocks = text_blocks(text)
        self.assertGreater(len(blocks), 1)
        self.assertTrue(all(len(b['text']) <= PARAGRAPH_SPLIT for b in blocks))
        self.assertEqual(sum(len(b['text'].split()) for b in blocks), len(text.split()))
