"""
Phase D backend: inline runs, upload imports, image assets, and the
edit_document / edit_deck operations.

What these pin:
* a spec with runs and alignment validates, renders to a real .docx (read
  back with python-docx) and to a PDF, and extracts with its markers;
* an uploaded Word file converts to an editable spec (headings, bold runs,
  lists, tables, images), keeping the original upload as version 1;
* an uploaded deck converts slide by slide to the nearest layout;
* an import refuses a file that is already editable, and a stale one;
* spec images are served only when the spec names them;
* block and slide ops apply, refuse bad indices, and convert uploads first.
"""
from __future__ import annotations

import io
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from inference import vfs
from inference.models import Document, DocumentVersion

User = get_user_model()

_MEDIA = tempfile.mkdtemp(prefix='import-')

def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new('RGB', (8, 8), (200, 30, 30)).save(buf, format='PNG')
    return buf.getvalue()


PNG = _png()


def _uploaded_docx() -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading('Quarterly Report', level=1)
    para = doc.add_paragraph()
    para.add_run('Revenue grew ')
    bold = para.add_run('38%')
    bold.bold = True
    doc.add_paragraph('First point', style='List Bullet')
    doc.add_paragraph('Second point', style='List Bullet')
    doc.add_heading('Details', level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'a'
    table.cell(0, 1).text = 'b'
    table.cell(1, 0).text = '1'
    table.cell(1, 1).text = '2'
    doc.add_picture(io.BytesIO(PNG), width=docx.shared.Inches(1))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _uploaded_deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = 'Kickoff'
    slide.placeholders[1].text = 'Welcome aboard'
    bullets_layout = prs.slide_layouts[1]
    second = prs.slides.add_slide(bullets_layout)
    second.shapes.title.text = 'Goals'
    second.placeholders[1].text_frame.text = 'Ship it'
    second.placeholders[1].text_frame.paragraphs[0].level = 0
    blank = prs.slides.add_slide(prs.slide_layouts[6])
    blank.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(1), Inches(2), Inches(2))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@override_settings(MEDIA_ROOT=_MEDIA)
class ImportCase(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(_MEDIA, ignore_errors=True)

    def setUp(self):
        self.owner = User.objects.create_user('owner', 'o@example.com', 'pw')
        self.me = APIClient()
        self.me.force_authenticate(self.owner)
        self.scope = vfs.build_scope(self.owner, 'full')

    def store(self, name, file_type, data: bytes):
        doc = Document(user=self.owner, name=name, file_type=file_type,
                       file_size=len(data), content_text='', status='stored', metadata={})
        doc.file.save(name, ContentFile(data), save=False)
        doc.save()
        return doc


class RunsTests(TestCase):
    def test_runs_and_alignment_validate_and_render(self):
        import docx

        from chat.tools.office import document, pdf

        spec = document.validate({
            'title': 'Runs', 'blocks': [
                {'type': 'heading', 'level': 2, 'align': 'center',
                 'runs': [{'text': 'Hi ', 'bold': True}, {'text': 'there', 'italic': True}]},
                {'type': 'paragraph', 'align': 'justify',
                 'runs': [{'text': 'Visit '}, {'text': 'us', 'link': 'https://example.com'}]},
                {'type': 'bullets', 'items': [
                    {'runs': [{'text': 'done', 'strike': True}]},
                    'plain **bold**',
                ]},
            ]})
        data = document.render(spec, {})
        read = docx.Document(io.BytesIO(data))
        heading = read.paragraphs[1]
        self.assertEqual(heading.runs[0].bold, True)
        self.assertEqual(heading.runs[1].italic, True)
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        self.assertEqual(heading.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        body = read.paragraphs[2].text
        self.assertEqual(body, 'Visit us')
        out = pdf.render(spec, {})
        self.assertTrue(out.startswith(b'%PDF'))
        text = document.extract_text(spec)
        self.assertIn('**Hi **', text)
        self.assertIn('[us](https://example.com)', text)

    def test_marker_text_still_validates(self):
        from chat.tools.office import document

        spec = document.validate({'title': 't', 'blocks': [
            {'type': 'paragraph', 'text': 'a **b** word'}]})
        self.assertEqual(spec['blocks'][0]['text'], 'a **b** word')


class DocxImportTests(ImportCase):
    def test_an_upload_converts_and_keeps_version_one(self):
        from chat.tools.office import document

        data = _uploaded_docx()
        doc = self.store('Report.docx', 'docx', data)
        resp = self.me.post(f'/api/inference/documents/{doc.id}/import/')
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertTrue(body['converted'])
        stored = Document.objects.get(id=doc.id)
        spec = stored.metadata['spec']
        validated = document.validate({'title': spec['title'], 'subtitle': spec.get('subtitle', ''),
                                       'theme': spec.get('theme', 'clean'),
                                       'blocks': spec['blocks']})
        kinds = [b['type'] for b in validated['blocks']]
        self.assertIn('heading', kinds)
        self.assertIn('bullets', kinds)
        self.assertIn('table', kinds)
        self.assertIn('image', kinds)
        text = document.extract_text(validated)
        self.assertIn('**38%**', text)
        # The original upload is version 1.
        version = DocumentVersion.objects.get(document_id=doc.id)
        with version.file.open('rb') as fh:
            self.assertEqual(fh.read(), data)

    def test_a_rename_does_not_make_an_open_editor_stale(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        before = Document.objects.get(id=doc.id).updated_at
        resp = self.me.patch(f'/api/inference/documents/{doc.id}/',
                             {'name': 'Renamed.docx'}, format='json')
        self.assertEqual(resp.status_code, 200, resp.content)
        after = Document.objects.get(id=doc.id)
        self.assertEqual(after.name, 'Renamed.docx')
        self.assertEqual(after.updated_at, before)
        resp = self.me.post(f'/api/inference/documents/{doc.id}/import/',
                            {'expected_updated_at': before.isoformat()}, format='json')
        self.assertEqual(resp.status_code, 200, resp.content)

    def test_an_editable_file_needs_no_import(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        self.me.post(f'/api/inference/documents/{doc.id}/import/')
        resp = self.me.post(f'/api/inference/documents/{doc.id}/import/')
        self.assertEqual(resp.status_code, 400)

    def test_a_stale_import_is_refused(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        before = Document.objects.get(id=doc.id).updated_at.isoformat()
        # A write after the caller opened it (a rename would not do: it
        # keeps `updated_at`, the etag, on purpose).
        Document.objects.get(id=doc.id).save()
        resp = self.me.post(f'/api/inference/documents/{doc.id}/import/',
                            {'expected_updated_at': before}, format='json')
        self.assertEqual(resp.status_code, 412)


class DeckImportTests(ImportCase):
    def test_slides_map_to_the_nearest_layout(self):
        data = _uploaded_deck()
        doc = self.store('Kickoff.pptx', 'pptx', data)
        resp = self.me.post(f'/api/inference/documents/{doc.id}/import/')
        self.assertEqual(resp.status_code, 200, resp.content)
        stored = Document.objects.get(id=doc.id)
        slides = stored.metadata['spec']['slides']
        layouts = [s['layout'] for s in slides]
        self.assertEqual(layouts[0], 'title')
        self.assertEqual(slides[0]['subtitle'], 'Welcome aboard')
        self.assertEqual(layouts[1], 'bullets')
        self.assertEqual(layouts[2], 'image')
        self.assertEqual(DocumentVersion.objects.filter(document_id=doc.id).count(), 1)


class AssetTests(ImportCase):
    def test_named_images_serve_and_strangers_404(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        self.me.post(f'/api/inference/documents/{doc.id}/import/')
        stored = Document.objects.get(id=doc.id)
        images = [b['path'] for b in stored.metadata['spec']['blocks']
                  if b.get('type') == 'image']
        self.assertTrue(images)
        resp = self.me.get(f'/api/inference/documents/{stored.id}/asset/', {'path': images[0]})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'image/png')
        body = b''.join(resp.streaming_content)
        self.assertTrue(body.startswith(b'\x89PNG'))
        resp = self.me.get(f'/api/inference/documents/{stored.id}/asset/', {'path': '/nope.png'})
        self.assertEqual(resp.status_code, 404)

    def test_a_non_raster_path_is_never_served(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        self.store('evil.svg', 'image', b'<svg xmlns="http://www.w3.org/2000/svg">'
                                        b'<script>alert(1)</script></svg>')
        Document.objects.filter(id=doc.id).update(metadata={'spec': {
            'title': 'R', 'blocks': [{'type': 'image', 'path': '/evil.svg', 'caption': ''}]}})
        resp = self.me.get(f'/api/inference/documents/{doc.id}/asset/', {'path': '/evil.svg'})
        self.assertEqual(resp.status_code, 404)

    def test_insert_image_upload_beside_the_document(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        upload = SimpleUploadedFile('chart.png', PNG, content_type='image/png')
        resp = self.me.post(f'/api/inference/documents/{doc.id}/images/', {'file': upload})
        self.assertEqual(resp.status_code, 201, resp.content)
        self.assertTrue(resp.json()['path'].endswith('chart.png'))
        upload = SimpleUploadedFile('notes.txt', b'x', content_type='text/plain')
        resp = self.me.post(f'/api/inference/documents/{doc.id}/images/', {'file': upload})
        self.assertEqual(resp.status_code, 400)


class EditOpsTests(ImportCase):
    def test_document_ops_apply_and_refuse(self):
        doc = self.store('Report.docx', 'docx', _uploaded_docx())
        out = vfs.edit_document(self.scope, '/Report.docx', [
            {'op': 'insert', 'index': 0,
             'blocks': [{'type': 'paragraph', 'text': 'Inserted'}]},
            {'op': 'find_replace', 'find': 'Revenue', 'replace': 'Income'},
        ])
        self.assertIn('inserted 1 block(s) at 0', out['applied'])
        self.assertTrue(out['imported_upload'])
        stored = Document.objects.get(id=doc.id)
        self.assertEqual(stored.metadata['spec']['blocks'][0]['text'], 'Inserted')
        self.assertIn('Income', stored.content_text)
        with self.assertRaises(vfs.VfsError):
            vfs.edit_document(self.scope, '/Report.docx', [
                {'op': 'delete', 'index': 99}])
        with self.assertRaises(vfs.VfsError):
            vfs.edit_document(self.scope, '/Report.docx', [
                {'op': 'find_replace', 'find': 'no such words'}])

    def test_deck_ops_apply_and_refuse(self):
        doc = self.store('Kickoff.pptx', 'pptx', _uploaded_deck())
        out = vfs.edit_deck(self.scope, '/Kickoff.pptx', [
            {'op': 'duplicate', 'index': 0},
            {'op': 'set', 'index': 1, 'fields': {'subtitle': 'Edited'}},
            {'op': 'move', 'index': 1, 'to': 0},
            {'op': 'remove', 'index': 3},
        ])
        self.assertEqual(len(out['applied']), 4)
        stored = Document.objects.get(id=doc.id)
        slides = stored.metadata['spec']['slides']
        self.assertEqual(len(slides), 3)
        self.assertEqual(slides[0]['subtitle'], 'Edited')
        with self.assertRaises(vfs.VfsError):
            vfs.edit_deck(self.scope, '/Kickoff.pptx', [{'op': 'remove', 'index': 9}])

    def test_the_agent_tool_edits_through_files(self):
        import json

        from asgiref.sync import async_to_sync

        from chat.tools import files as file_tools

        self.store('Report.docx', 'docx', _uploaded_docx())
        payload = json.loads(async_to_sync(file_tools.edit_document)(
            {'path': '/Report.docx',
             'ops': [{'op': 'insert', 'index': 0,
                      'blocks': [{'type': 'paragraph', 'text': 'Hi'}]}]},
            {'file_scope': self.scope}))
        self.assertIn('inserted', str(payload.get('applied')))


class ImportOrderTests(TestCase):
    """Review fixes (2026-09-25)."""

    def test_tables_stay_where_they_were(self):
        import io

        from docx import Document as Docx

        from inference import importers

        d = Docx()
        d.add_paragraph('Before the table')
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Head'
        table.cell(1, 0).text = 'value'
        d.add_paragraph('After the table')
        buf = io.BytesIO()
        d.save(buf)
        spec, _, _ = importers.docx_to_spec(buf.getvalue(), 'Report')
        self.assertEqual([b['type'] for b in spec['blocks']],
                         ['paragraph', 'table', 'paragraph'])

    def test_a_deck_is_named_for_its_first_title_or_the_file(self):
        import io

        from pptx import Presentation

        from inference import importers

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 10, 10) \
            .text_frame.text = 'no title here'
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = 'Last slide'
        buf = io.BytesIO()
        prs.save(buf)
        spec, _, _ = importers.pptx_to_spec(buf.getvalue(), 'Quarterly')
        self.assertEqual(spec['title'], 'Quarterly')
