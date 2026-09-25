"""
The shared formula evaluator (`inference/formulas.py`).

Checked as a table of expected results — one workbook, every function — plus
the properties that make it safe to run model-written strings: anything
outside the language is refused rather than run, cycles fail loudly, and the
benchmark reads through the same implementation (`eval/office_files.py`
re-exports this module).
"""
from __future__ import annotations

import datetime

from django.test import SimpleTestCase

from inference import formulas


def _book():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    for c, h in enumerate(['item', 'qty', 'price'], start=1):
        ws.cell(row=1, column=c, value=h)
    for r, row in enumerate([['apple', 3, 1.5], ['pear', 2, 2.5], ['apple', 1, 1.5]], start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    other = wb.create_sheet('Meta')
    other['A1'] = 10
    return wb


class EvaluateTests(SimpleTestCase):
    def setUp(self):
        self.wb = _book()

    def check(self, formula, expected):
        self.assertEqual(formulas.evaluate(self.wb, 'Data', formula), expected, formula)

    def test_arithmetic_and_references(self):
        self.check('=B2+C2', 4.5)
        self.check('=B2*B3+C4^2', 3 * 2 + 1.5 ** 2)
        self.check('=SUM(B2:B4)', 6)
        self.check('=AVERAGE(B2:B4)', 2.0)
        self.check('=MIN(C2:C4)', 1.5)
        self.check('=MAX(C2:C4)', 2.5)
        self.check('=COUNT(B2:B4)', 3.0)
        self.check('=COUNTA(A2:A4)', 3.0)
        self.check('=ROUND(C3,1)', 2.5)
        self.check('=ABS(B2-10)', 7)
        self.check('=Meta!A1*2', 20)
        self.check('=SUM(Data!B2:B4)', 6)

    def test_comparisons_and_logic(self):
        self.check('=B2=3', True)
        self.check('=B2<>3', False)
        self.check('=B2>2', True)
        self.check('=A2="APPLE"', True)
        self.check('=IF(B2>2,"big","small")', 'big')
        self.check('=IF(B3>2,"big","small")', 'small')
        self.check('=IFERROR(1/0,"caught")', 'caught')
        self.check('=IFERROR(B2,"caught")', 3)
        self.check('=AND(B2=3,C2=1.5)', True)
        self.check('=AND(B2=3,C2=9)', False)
        self.check('=OR(B2=9,C2=1.5)', True)
        self.check('=NOT(B2=9)', True)

    def test_text(self):
        self.check('=CONCAT(A2," x",B2)', 'apple x3')
        self.check('=A2&"!"', 'apple!')
        self.check('=LEN(A2)', 5)
        self.check('=UPPER(A2)', 'APPLE')
        self.check('=LOWER("Pear")', 'pear')
        self.check('=TRIM("  a  b  ")', 'a b')

    def test_conditionals(self):
        self.check('=SUMIF(A2:A4,"apple",B2:B4)', 4.0)
        self.check('=SUMIF(B2:B4,">2")', 3.0)
        self.check('=COUNTIF(A2:A4,"apple")', 2.0)
        self.check('=COUNTIF(A2:A4,"a*")', 2.0)
        self.check('=AVERAGEIF(A2:A4,"apple",C2:C4)', 1.5)

    def test_lookups(self):
        self.check('=VLOOKUP("pear",A2:C4,3)', 2.5)
        self.check('=VLOOKUP("pear",A2:C4,2)', 2)
        self.check('=XLOOKUP("pear",A2:A4,C2:C4)', 2.5)
        self.check('=XLOOKUP("plum",A2:A4,C2:C4,"none")', 'none')

    def test_today(self):
        self.assertEqual(formulas.evaluate(self.wb, 'Data', '=TODAY()'), datetime.date.today())

    def test_lazy_branches_never_run(self):
        # Eager evaluation would divide by zero before IF chose its branch.
        self.check('=IF(1>2,1/0,"safe")', 'safe')
        self.check('=IFERROR(SUM(B2:B4),"no")', 6)

    def test_cycles_fail_loudly(self):
        self.wb['Data']['D2'] = '=E2'
        self.wb['Data']['E2'] = '=D2'
        with self.assertRaises(formulas.FormulaError):
            formulas.cell(self.wb, 'Data', 'D2')

    def test_outside_the_language_is_refused(self):
        for formula in ('=__import__("os").getcwd()', '=A1.__class__',
                        '=POWER(B2,2)', '=[x for x in B2:B4]', '=B2()',
                        '=SUM(B2:B4,', '=NOSUCHFN(B2)'):
            with self.subTest(formula=formula):
                with self.assertRaises(formulas.FormulaError):
                    formulas.evaluate(self.wb, 'Data', formula)

    def test_missing_sheet_and_no_match(self):
        with self.assertRaises(formulas.FormulaError):
            formulas.evaluate(self.wb, 'Data', '=SUM(Gone!A1:A2)')
        with self.assertRaises(formulas.FormulaError):
            formulas.evaluate(self.wb, 'Data', '=VLOOKUP("plum",A2:C4,2)')
        with self.assertRaises(formulas.FormulaError):
            formulas.evaluate(self.wb, 'Data', '=AVERAGEIF(A2:A4,"plum",C2:C4)')
        with self.assertRaises(formulas.FormulaError):
            formulas.evaluate(self.wb, 'Data', '=1/0')


class DelegationTests(SimpleTestCase):
    def test_the_benchmark_reads_through_this_module(self):
        from eval import office_files

        self.assertIs(office_files.FormulaError, formulas.FormulaError)
        wb = office_files.workbook(_book_bytes())
        self.assertEqual(office_files.cell(wb, 'Data', 'B2'), 3)


def _book_bytes() -> bytes:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws['A1'] = 'item'
    ws['B2'] = 3
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ScaleTests(SimpleTestCase):
    """Review fixes (2026-09-25): long chains and huge ranges."""

    def test_a_long_running_total_calculates(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'S'
        for r in range(1, 501):
            ws[f'A{r}'] = 1
            ws[f'B{r}'] = '=A1' if r == 1 else f'=B{r - 1}+A{r}'
        self.assertEqual(formulas.cell(wb, 'S', 'B500'), 500)

    def test_a_range_past_the_data_creates_no_cells(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'S'
        for r in range(1, 11):
            ws[f'A{r}'] = r
        ws['AB1'] = '=SUM(A1:Z1000000)'
        self.assertEqual(formulas.cell(wb, 'S', 'AB1'), 55)
        self.assertLess(len(ws._cells), 20)

    def test_too_much_real_data_is_refused(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'S'
        ws.cell(row=10_000, column=26, value=1)  # the used area is 260k cells
        with self.assertRaises(formulas.FormulaError):
            formulas.evaluate(wb, 'S', '=SUM(A1:Z10000)')

    def test_blank_counting_stays_exact_for_a_normal_range(self):
        self.assertEqual(
            formulas.evaluate(_book(), 'Data',
                              '=COUNTIF(A1:A10,"")'),
            6,
        )
