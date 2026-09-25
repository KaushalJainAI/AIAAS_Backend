"""
Sheets on Univer, phase C backend (`inference/sheets.py`, `formulas.py`,
`chat/tools/office/edit.py` and the values plumbing).

What these pin:
* a snapshot round trip keeps values, formulas, styles, merges, dimensions,
  freeze panes, sheet order and names — and the file still opens in Excel;
* a renamed sheet is renamed, not duplicated next to an empty new one;
* charts survive a snapshot round trip and a cell edit (openpyxl carries what
  it does not understand through a load/save untouched);
* inserting or deleting rows/columns mid-sheet shifts formula references and
  merges; references into a deleted band become #REF!, which the evaluator
  refuses loudly;
* format / freeze / widths paint the file;
* CSV export and the office grid show calculated values, not formula text.
"""
from __future__ import annotations

import io
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from inference import formulas, sheets
from inference.models import Document

User = get_user_model()

_MEDIA = tempfile.mkdtemp(prefix='sheets-')


def _styled_book():
    import datetime

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws['A1'] = 'item'
    ws['B1'] = 'qty'
    ws['A2'] = 'apple'
    ws['B2'] = 3
    ws['A3'] = 'pear'
    ws['B3'] = 2
    ws['B4'] = '=SUM(B2:B3)'
    ws['C2'] = datetime.datetime(2024, 1, 15, 12, 0)
    ws['C2'].number_format = 'yyyy-mm-dd'
    bold_red = Font(bold=True, color='FFFF0000', size=14, name='Calibri')
    ws['A1'].font = bold_red
    ws['B1'].font = bold_red
    ws['A2'].fill = PatternFill(patternType='solid', fgColor='FFFFFF00')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws['B2'].number_format = '#,##0.00'
    thin = Side(style='thin', color='FF2A78D6')
    ws['A1'].border = Border(top=thin, right=thin, bottom=thin, left=thin)
    ws.merge_cells('A5:B5')
    ws['A5'] = 'merged note'
    ws.column_dimensions['A'].width = 20
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = 'B2'
    second = wb.create_sheet('Meta')
    second['A1'] = '=Data!B4*2'
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _open(data: bytes):
    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(data))


@override_settings(MEDIA_ROOT=_MEDIA)
class SnapshotTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(_MEDIA, ignore_errors=True)

    def test_a_round_trip_keeps_everything_and_stays_openable(self):
        data = _styled_book()
        snap = sheets.to_snapshot(data)
        self.assertEqual(snap['appVersion'], sheets.APP_VERSION)
        self.assertEqual(snap['sheetOrder'], ['sheet-0', 'sheet-1'])
        out = sheets.apply_snapshot(data, snap)
        # Byte-compatible with Excel: opens again, same structure.
        again = sheets.to_snapshot(out)
        self.assertEqual([s['name'] for s in again['sheets'].values()], ['Data', 'Meta'])

        wb = _open(out)
        try:
            ws = wb['Data']
            self.assertEqual(ws['B4'].value, '=SUM(B2:B3)')
            self.assertEqual(ws['A1'].font.bold, True)
            self.assertEqual(ws['A1'].font.color.rgb, 'FFFF0000')
            self.assertEqual(ws['A2'].fill.fgColor.rgb, 'FFFFFF00')
            self.assertEqual(ws['A2'].alignment.horizontal, 'center')
            self.assertEqual(ws['A2'].alignment.wrap_text, True)
            self.assertEqual(ws['B2'].number_format, '#,##0.00')
            self.assertEqual(ws['A1'].border.top.style, 'thin')
            self.assertEqual(ws['C2'].number_format, 'yyyy-mm-dd')
            self.assertEqual([str(r) for r in ws.merged_cells.ranges], ['A5:B5'])
            self.assertAlmostEqual(ws.column_dimensions['A'].width, 20, places=1)
            self.assertAlmostEqual(ws.row_dimensions[2].height, 30, places=1)
            self.assertEqual(ws.freeze_panes, 'B2')
            self.assertEqual(wb['Meta']['A1'].value, '=Data!B4*2')
        finally:
            wb.close()

    def test_values_travel_as_numbers_and_dates_as_serials(self):
        snap = sheets.to_snapshot(_styled_book())
        cells = snap['sheets']['sheet-0']['cellData']
        self.assertEqual(cells[1][1]['v'], 3)
        self.assertIn('n', snap['styles'][cells[1][1]['s']])
        # Formulas carry the backend-calculated value beside the source, so
        # the grid shows a number before the formula worker computes.
        self.assertEqual(cells[3][1]['f'], '=SUM(B2:B3)')
        self.assertEqual(cells[3][1]['v'], 5.0)
        serial = cells[1][2]['v']
        self.assertIsInstance(serial, float)
        self.assertIn('yyyy-mm-dd', str(snap['styles'][cells[1][2]['s']]['n']))

    def test_a_renamed_sheet_is_renamed_not_duplicated(self):
        snap = sheets.to_snapshot(_styled_book())
        snap['sheets']['sheet-0']['name'] = 'Renamed'
        out = sheets.apply_snapshot(_styled_book(), snap)
        wb = _open(out)
        try:
            self.assertEqual(wb.sheetnames, ['Renamed', 'Meta'])
            self.assertEqual(wb['Renamed']['B2'].value, 3)
        finally:
            wb.close()

    def test_an_emptied_cell_is_cleared_not_kept(self):
        snap = sheets.to_snapshot(_styled_book())
        del snap['sheets']['sheet-0']['cellData'][1][1]  # B2 emptied in the app
        out = sheets.apply_snapshot(_styled_book(), snap)
        wb = _open(out)
        try:
            self.assertIsNone(wb['Data']['B2'].value)
        finally:
            wb.close()

    def test_a_bad_snapshot_is_refused_not_applied(self):
        with self.assertRaises(sheets.SnapshotError):
            sheets.apply_snapshot(_styled_book(), {'sheets': [1, 2]})

    def test_charts_survive_a_snapshot_round_trip(self):
        from chat.tools.office import workbook

        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'm'}, {'header': 'v', 'type': 'number'}],
            'rows': [['a', 1], ['b', 2]],
            'chart': {'kind': 'bar', 'title': 't', 'x': 'm', 'y': 'v'},
        }]})
        data, _ = workbook.render(spec)
        out = sheets.apply_snapshot(data, sheets.to_snapshot(data))
        wb = _open(out)
        try:
            self.assertTrue(formulas.has_chart(wb))
        finally:
            wb.close()


class StructuralEditTests(TestCase):
    def test_insert_rows_shifts_formulas_and_merges(self):
        from chat.tools.office import edit

        data = _styled_book()
        change = edit.validate({'sheet': 'Data', 'insert_rows': {'row': 2, 'count': 1}})
        out, _ = edit.apply(data, change)
        wb = _open(out)
        try:
            # The total moved down with its rows and still points at them.
            self.assertEqual(wb['Data']['B5'].value, '=SUM(B3:B4)')
            self.assertEqual(wb['Data']['A3'].value, 'apple')
            self.assertEqual([str(r) for r in wb['Data'].merged_cells.ranges], ['A6:B6'])
            # ... while the other sheet's cross-sheet reference followed it.
            self.assertEqual(wb['Meta']['A1'].value, '=Data!B5*2')
        finally:
            wb.close()

    def test_delete_rows_makes_refs_ref_loudly(self):
        from chat.tools.office import edit

        data = _styled_book()
        change = edit.validate({'sheet': 'Data', 'delete_rows': {'row': 2, 'count': 1}})
        out, _ = edit.apply(data, change)
        wb = _open(out)
        try:
            self.assertIn('#REF!', wb['Data']['B3'].value)
            with self.assertRaises(formulas.FormulaError):
                formulas.cell(wb, 'Data', 'B3')
        finally:
            wb.close()

    def test_insert_and_delete_cols_shift(self):
        from chat.tools.office import edit

        data = _styled_book()
        out, _ = edit.apply(data, edit.validate({'sheet': 'Data', 'insert_cols': {'col': 'A'}}))
        wb = _open(out)
        try:
            self.assertEqual(wb['Data']['C4'].value, '=SUM(C2:C3)')
        finally:
            wb.close()
        out, _ = edit.apply(data, edit.validate({'sheet': 'Data', 'delete_cols': {'col': 'C'}}))
        wb = _open(out)
        try:
            # C gone: the date vanished and the total still sums B.
            self.assertIsNone(wb['Data']['C2'].value)
            self.assertEqual(wb['Data']['B4'].value, '=SUM(B2:B3)')
        finally:
            wb.close()

    def test_format_freeze_and_widths(self):
        from chat.tools.office import edit

        data = _styled_book()
        change = edit.validate({
            'sheet': 'Data',
            'format': {'range': 'A1:B1',
                       'style': {'bold': True, 'fill': '#2a78d6', 'alignment': 'center',
                                 'number_format': '#,##0', 'border': {'bottom': 'thin'}}},
            'freeze': 'C3',
            'widths': {'B': 30},
        })
        out, report = edit.apply(data, change)
        self.assertEqual(report['formatted'], 'A1:B1')
        self.assertEqual(report['frozen'], 'C3')
        wb = _open(out)
        try:
            self.assertEqual(wb['Data']['A1'].fill.fgColor.rgb, 'FF2A78D6')
            self.assertEqual(wb['Data']['A1'].alignment.horizontal, 'center')
            self.assertEqual(wb['Data']['B1'].number_format, '#,##0')
            self.assertEqual(wb['Data']['A1'].border.bottom.style, 'thin')
            self.assertEqual(wb['Data'].freeze_panes, 'C3')
            self.assertAlmostEqual(wb['Data'].column_dimensions['B'].width, 30, places=1)
        finally:
            wb.close()

    def test_charts_survive_a_cell_edit(self):
        from chat.tools.office import edit, workbook

        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'm'}, {'header': 'v', 'type': 'number'}],
            'rows': [['a', 1], ['b', 2]],
            'chart': {'kind': 'bar', 'title': 't', 'x': 'm', 'y': 'v'},
        }]})
        data, _ = workbook.render(spec)
        out, _ = edit.apply(data, edit.validate({'sheet': 'S', 'set_cells': [{'cell': 'B2', 'value': 5}]}))
        wb = _open(out)
        try:
            self.assertTrue(formulas.has_chart(wb))
            self.assertEqual(wb['S']['B2'].value, 5)
        finally:
            wb.close()

    def test_refusals(self):
        from chat.tools.office import edit
        from chat.tools.office.spec import SpecError

        with self.assertRaises(SpecError):
            edit.validate({'sheet': 'Data', 'insert_rows': {'row': 0}})
        with self.assertRaises(SpecError):
            edit.validate({'sheet': 'Data', 'format': {'range': 'ZZZ', 'style': {'bold': True}}})
        with self.assertRaises(SpecError):
            edit.validate({'sheet': 'Data', 'freeze': 'nowhere'})
        with self.assertRaises(SpecError):
            edit.validate({'sheet': 'Data', 'widths': {'A': -1}})
        with self.assertRaises(SpecError):
            edit.validate({'sheet': 'Data'})


@override_settings(MEDIA_ROOT=_MEDIA)
class CalculatedValueTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(_MEDIA, ignore_errors=True)

    def setUp(self):
        self.user = User.objects.create_user('owner', 'o@example.com', 'pw')

    def _document(self, data: bytes, name='B.xlsx'):
        doc = Document(user=self.user, name=name, file_type='xlsx',
                       file_size=len(data), status='stored', metadata={})
        doc.file.save(name, ContentFile(data), save=False)
        doc.save()
        return doc

    def test_csv_export_calculates_instead_of_quoting_formulas(self):
        from chat.tools.office import workbook
        from inference import export

        # xlsxwriter stores formulas without cached values: a reader asking
        # for values sees nothing, which is what the old export quoted.
        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'a', 'type': 'number'},
                                     {'header': 'b', 'type': 'number'}],
            'rows': [[2, 3]], 'totals': True,
        }]})
        data, _ = workbook.render(spec)
        doc = self._document(data)
        body = export.workbook_csv(doc).decode('utf-8')
        self.assertIn('2.0', body)
        self.assertIn('3.0', body)
        self.assertNotIn('SUM', body)

    def test_the_office_grid_carries_values_beside_formulas(self):
        from chat.tools.office import workbook
        from inference import office_edit

        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'a', 'type': 'number'},
                                     {'header': 'b', 'type': 'number'}],
            'rows': [[2, 3]], 'totals': True,
        }]})
        data, _ = workbook.render(spec)
        doc = self._document(data)
        grid = office_edit.workbook_grid(doc)
        sheet = grid['sheets'][0]
        total_row = sheet['rows'][-1]
        total_values = sheet['values'][-1]
        self.assertTrue(any(isinstance(v, str) and v.startswith('=') for v in total_row))
        self.assertEqual([v for v in total_values if isinstance(v, (int, float))], [2.0, 3.0])

    def test_read_workbook_reports_values_and_formulas(self):
        from chat.tools.office import workbook
        from chat.tools import office as office_tools

        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'a', 'type': 'number'},
                                     {'header': 'b', 'type': 'number'}],
            'rows': [[2, 3]], 'totals': True,
        }]})
        data, _ = workbook.render(spec)
        result = office_tools._read_sheet(data, None, None)
        total_values = result['values'][-1]
        total_formulas = result['formulas'][-1]
        self.assertEqual([v for v in total_values if isinstance(v, (int, float))], [2.0, 3.0])
        self.assertTrue(any(isinstance(f, str) and f.startswith('=') for f in total_formulas))

    def test_a_snapshot_draft_renders(self):
        from chat.tools.office import workbook
        from inference import drafts, office_edit

        spec = workbook.validate({'sheets': [{
            'name': 'S', 'columns': [{'header': 'a', 'type': 'number'}],
            'rows': [[1]],
        }]})
        data, _ = workbook.render(spec)
        doc = self._document(data)
        snap = sheets.to_snapshot(data)
        first_id = snap['sheetOrder'][0]
        snap['sheets'][first_id]['cellData'][1][0] = {'v': 9}
        doc, _stamp = drafts.save_draft(doc, {'snapshot': snap}, None)
        self.assertEqual(doc.metadata['draft']['kind'], 'snapshot')
        self.assertTrue(drafts.maybe_render_draft(doc.id, force=True))
        grid = office_edit.workbook_grid(Document.objects.get(id=doc.id))
        self.assertEqual(grid['sheets'][0]['rows'][1][0], 9)


class SnapshotStructureTests(TestCase):
    """Review fixes (2026-09-25): deletes, order and filled formulas."""

    def test_a_deleted_sheet_is_removed_and_tab_order_is_kept(self):
        data = _styled_book()
        snap = sheets.to_snapshot(data)
        data_id, meta_id = snap['sheetOrder']
        snap['sheetOrder'] = [meta_id]
        del snap['sheets'][data_id]
        # Meta's formula names the deleted sheet; that is the user's call.
        wb = _open(sheets.apply_snapshot(data, snap))
        try:
            self.assertEqual(wb.sheetnames, ['Meta'])
        finally:
            wb.close()

    def test_tabs_follow_the_snapshot_order(self):
        data = _styled_book()
        snap = sheets.to_snapshot(data)
        snap['sheetOrder'] = list(reversed(snap['sheetOrder']))
        wb = _open(sheets.apply_snapshot(data, snap))
        try:
            self.assertEqual(wb.sheetnames, ['Meta', 'Data'])
        finally:
            wb.close()

    def test_shared_formulas_are_saved_as_formulas(self):
        data = _styled_book()
        snap = sheets.to_snapshot(data)
        data_id = snap['sheetOrder'][0]
        cells = snap['sheets'][data_id]['cellData']
        # How Univer stores a drag-filled formula: `f` + `si` on the first
        # cell, `si` (and a value) on the rest.
        cells.setdefault('1', {})['3'] = {'f': '=B2*2', 'si': 'fill1', 'v': 6}
        cells.setdefault('2', {})['3'] = {'si': 'fill1', 'v': 4}
        wb = _open(sheets.apply_snapshot(data, snap))
        try:
            self.assertEqual(wb['Data']['D2'].value, '=B2*2')
            self.assertEqual(wb['Data']['D3'].value, '=B3*2')
        finally:
            wb.close()
